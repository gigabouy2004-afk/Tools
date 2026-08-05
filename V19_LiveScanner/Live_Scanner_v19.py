#!/usr/bin/env python3
"""V19 deterministic D-1 momentum qualification engine.

V19 is intentionally narrow.  It evaluates only completed daily candles before
the execution date, applies the owner-approved EMA/MACD/return qualification,
and calculates stochastic timing only for qualified securities.

Exit codes:
    0 - complete run and workbook written
    2 - command-line or pre-flight validation error
    3 - provider-degraded partial workbook written
    4 - invalid run workbook written (no security could be processed)
    5 - output workbook could not be written
"""

from __future__ import annotations

import argparse
import math
import os
import random
import re
import sys
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Callable, Iterable, Sequence

import numpy as np
import openpyxl
import pandas as pd
import yfinance as yf
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from yfinance.exceptions import YFPricesMissingError, YFRateLimitError, YFTickerMissingError


ENGINE_NAME = "V19 D-1 Deterministic Momentum Qualification Engine"
ENGINE_VERSION = "19.0.0"
SCRIPT_PATH = Path(__file__).resolve()

SHORT_WINDOW_SESSIONS = 21
MEDIUM_WINDOW_SESSIONS = 63
LONG_WINDOW_SESSIONS = 126
EMA_MEDIUM_SPAN = 50
EMA_LONG_SPAN = 200
MACD_FAST_SPAN = 12
MACD_SLOW_SPAN = 26
MACD_SIGNAL_SPAN = 9
TIER_1_THRESHOLD = 0.20
TIER_2_THRESHOLD = 0.10

STOCH_LOOKBACK = 14
STOCH_K_SMOOTH = 3
STOCH_D_SMOOTH = 3
STOCH_OVERSOLD_LEVEL = 20.0
STOCH_RECENT_LOOKBACK = 14
DEFAULT_STOCH_ZONE_LOW = 45.0
DEFAULT_STOCH_ZONE_HIGH = 55.0

MIN_REQUIRED_SESSIONS = max(EMA_LONG_SPAN, LONG_WINDOW_SESSIONS + 1)
DEFAULT_HISTORY_LOOKBACK_DAYS = 1_100
DEFAULT_WORKERS = 4
DEFAULT_MAX_ATTEMPTS = 3
DEFAULT_REQUEST_TIMEOUT = 15.0
DEFAULT_REQUEST_INTERVAL = 0.35
DEFAULT_RETRY_BASE_SECONDS = 1.5
DEFAULT_PROGRESS_EVERY = 25

STATUS_TIER_1 = "QUALIFIED_TIER_1"
STATUS_TIER_2 = "QUALIFIED_TIER_2"
STATUS_REJECT = "REJECT"
STATUS_INVALID = "INVALID_DATA"

RUN_COMPLETE = "COMPLETE"
RUN_PARTIAL = "PARTIAL_PROVIDER_DEGRADED"
RUN_INVALID = "INVALID_RUN"

SHEET_NAMES = ["Summary", "Review", "Baseline Info", "Technical Data", "Scoring"]
SUPPORTED_INPUT_SUFFIXES = {".csv", ".txt", ".xlsx"}
HEADER_WORDS = {"SYMBOL", "TICKER", "CODE", "STOCK", "STOCK CODE", "STOCK_CODE"}
SYMBOL_PATTERN = re.compile(r"^[A-Z0-9][A-Z0-9.\-^=&]{0,39}$")


BASELINE_COLUMNS = [
    "Symbol",
    "Baseline_Start_Date",
    "Baseline_Start_Price",
    "D1_Date",
    "D1_Close",
    "D2_Date",
    "Total_Sessions_Available",
    "Execution_Date_D",
    "Calendar_Days_D1_to_D",
]

TECHNICAL_COLUMNS = [
    "Symbol",
    "Execution_Date_D",
    "D1_Date",
    "D2_Date",
    "D1_Close",
    "EMA_50",
    "EMA_200",
    "Price_vs_EMA50_%",
    "Price_vs_EMA200_%",
    "Price_Above_EMA50",
    "Price_Above_EMA200",
    "Bull_Zone_Status",
    "MACD_Line",
    "MACD_Signal",
    "MACD_Histogram",
    "MACD_Line_Above_Zero",
    "MACD_Signal_Above_Zero",
    "MACD_Line_Above_Signal",
    "MACD_Buyer_Zone",
    "Short_Window_Sessions",
    "Return_Short",
    "Medium_Window_Sessions",
    "Return_Medium",
    "Long_Window_Sessions",
    "Return_Long",
    "Momentum_Return_Band",
    "Qualification_Status",
    "Qualification_Tier",
    "Rejection_Reasons",
    "Stochastic_Calculated",
    "Stochastic_K_D1",
    "Stochastic_D_D1",
    "Stochastic_K_D2",
    "Stochastic_D_D2",
    "Stochastic_Crossover_Midpoint",
    "Stochastic_Fresh_Bullish_Crossover",
    "Stochastic_Near_50_Zone",
    "Stochastic_Recent_Oversold",
    "Stochastic_Status",
    "Data_Integrity_Flag",
]

SCORING_COLUMNS = [
    "Symbol",
    "Last Traded Price",
    "D1_Date",
    "Primary_Evaluation_Window",
    "Return_Short",
    "Return_Medium",
    "Return_Long",
    "Qualification_Tier",
    "Bull_Zone_Status",
    "MACD_Buyer_Zone",
    "Stochastic_K_D1",
    "Stochastic_D_D1",
    "Stochastic_K_D2",
    "Stochastic_D_D2",
    "Stochastic_Fresh_Bullish_Crossover",
    "Stochastic_Near_50_Zone",
    "Stochastic_Recent_Oversold",
    "Stochastic_Status",
]

REVIEW_COLUMNS = [
    "Symbol",
    "Status",
    "Qualification_Tier",
    "D1_Date",
    "Bull_Zone_Status",
    "MACD_Buyer_Zone",
    "Momentum_Return_Band",
    "Stochastic_Status",
    "Attempts",
    "Provider_Failure",
    "Error_Code",
    "Message",
]


class PreflightError(ValueError):
    """Raised before any market-data request is allowed to start."""


class DataExtractionError(RuntimeError):
    """A terminal per-symbol extraction or calculation failure."""

    def __init__(
        self,
        code: str,
        message: str,
        *,
        provider_failure: bool = False,
        retriable: bool = False,
    ) -> None:
        super().__init__(message)
        self.code = code
        self.provider_failure = provider_failure
        self.retriable = retriable


@dataclass(frozen=True)
class RuntimeConfig:
    execution_date: date
    explicit_execution_date: bool
    input_source: str
    raw_input_count: int
    duplicate_count: int
    output_path: Path
    overwrite: bool
    workers: int
    max_attempts: int
    request_timeout: float
    request_interval: float
    retry_base_seconds: float
    history_lookback_days: int
    progress_every: int
    stoch_zone_low: float
    stoch_zone_high: float

    @property
    def run_mode(self) -> str:
        return "BACKTEST_D1" if self.explicit_execution_date else "LIVE_D1"


@dataclass
class SymbolResult:
    index: int
    symbol: str
    status: str
    tier: str | None
    attempts: int
    provider_failure: bool
    error_code: str | None
    review: dict
    baseline: dict | None = None
    technical: dict | None = None
    scoring: dict | None = None


class RetryStatistics:
    """Thread-safe run diagnostics."""

    def __init__(self) -> None:
        self._lock = threading.Lock()
        self.request_attempts = 0
        self.retries = 0
        self.rate_limit_events = 0

    def record_attempt(self) -> None:
        with self._lock:
            self.request_attempts += 1

    def record_retry(self) -> None:
        with self._lock:
            self.retries += 1

    def record_rate_limit(self) -> None:
        with self._lock:
            self.rate_limit_events += 1


class ProviderGuard:
    """Global pacing and circuit pause shared by all Yahoo worker threads."""

    def __init__(self, minimum_interval: float, sleep_fn: Callable[[float], None] = time.sleep) -> None:
        self.minimum_interval = max(0.0, minimum_interval)
        self.sleep_fn = sleep_fn
        self._lock = threading.Lock()
        self._next_request_at = 0.0
        self._blocked_until = 0.0

    def wait_for_request_slot(self) -> None:
        while True:
            with self._lock:
                now = time.monotonic()
                available_at = max(self._next_request_at, self._blocked_until)
                wait_seconds = available_at - now
                if wait_seconds <= 0:
                    self._next_request_at = now + self.minimum_interval
                    return
            self.sleep_fn(min(wait_seconds, 1.0))

    def pause_after_rate_limit(self, seconds: float) -> None:
        with self._lock:
            self._blocked_until = max(self._blocked_until, time.monotonic() + max(0.0, seconds))


HistoryFetcher = Callable[[str, date, date, float], pd.DataFrame]


def build_argument_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description=ENGINE_NAME,
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument(
        "-c",
        "--codes",
        nargs="+",
        help="Stock codes; spaces and comma-separated segments are accepted.",
    )
    source.add_argument(
        "-i",
        "--input-file",
        help="CSV, TXT or XLSX file whose first column contains stock codes.",
    )
    parser.add_argument(
        "-d",
        "--execution-date",
        "--backtest-date",
        dest="execution_date",
        help="Decision date D in YYYY-MM-DD form. Candles dated D are excluded.",
    )
    parser.add_argument("-o", "--output", help="Output XLSX file path.")
    parser.add_argument("--overwrite", action="store_true", help="Permit replacement of an existing output file.")
    parser.add_argument("--workers", type=int, default=DEFAULT_WORKERS, help="Bounded Yahoo worker count.")
    parser.add_argument("--max-attempts", type=int, default=DEFAULT_MAX_ATTEMPTS)
    parser.add_argument("--request-timeout", type=float, default=DEFAULT_REQUEST_TIMEOUT)
    parser.add_argument("--request-interval", type=float, default=DEFAULT_REQUEST_INTERVAL)
    parser.add_argument("--retry-base-seconds", type=float, default=DEFAULT_RETRY_BASE_SECONDS)
    parser.add_argument("--history-lookback-days", type=int, default=DEFAULT_HISTORY_LOOKBACK_DAYS)
    parser.add_argument("--progress-every", type=int, default=DEFAULT_PROGRESS_EVERY)
    parser.add_argument("--stoch-zone-low", type=float, default=DEFAULT_STOCH_ZONE_LOW)
    parser.add_argument("--stoch-zone-high", type=float, default=DEFAULT_STOCH_ZONE_HIGH)
    return parser


def _split_code_values(values: Iterable[object]) -> list[str]:
    codes: list[str] = []
    for value in values:
        if value is None or (isinstance(value, float) and math.isnan(value)):
            continue
        for segment in str(value).replace(",", " ").split():
            candidate = segment.strip().upper()
            if candidate and candidate not in HEADER_WORDS:
                codes.append(candidate)
    return codes


def read_symbols_from_file(input_path: Path) -> list[str]:
    if not input_path.exists():
        raise PreflightError(f"Input file does not exist: {input_path}")
    if not input_path.is_file():
        raise PreflightError(f"Input path is not a file: {input_path}")
    suffix = input_path.suffix.lower()
    if suffix not in SUPPORTED_INPUT_SUFFIXES:
        allowed = ", ".join(sorted(SUPPORTED_INPUT_SUFFIXES))
        raise PreflightError(f"Unsupported input extension '{suffix}'. Allowed: {allowed}")
    try:
        if suffix == ".txt":
            raw_values = input_path.read_text(encoding="utf-8-sig").splitlines()
        elif suffix == ".csv":
            frame = pd.read_csv(input_path, header=None, dtype=str, keep_default_na=False)
            raw_values = frame.iloc[:, 0].tolist() if not frame.empty else []
        else:
            frame = pd.read_excel(input_path, sheet_name=0, header=None, dtype=str)
            raw_values = frame.iloc[:, 0].tolist() if not frame.empty else []
    except Exception as exc:
        raise PreflightError(f"Unable to read input file '{input_path}': {exc}") from exc
    return _split_code_values(raw_values)


def normalize_and_deduplicate_symbols(raw_codes: Sequence[str]) -> tuple[list[str], int]:
    invalid: list[str] = []
    unique: list[str] = []
    seen: set[str] = set()
    for raw in raw_codes:
        symbol = str(raw).strip().upper()
        if not SYMBOL_PATTERN.fullmatch(symbol):
            invalid.append(symbol or "<blank>")
            continue
        if symbol not in seen:
            seen.add(symbol)
            unique.append(symbol)
    if invalid:
        sample = ", ".join(invalid[:20])
        suffix = " ..." if len(invalid) > 20 else ""
        raise PreflightError(f"Invalid stock-code syntax: {sample}{suffix}")
    if not unique:
        raise PreflightError("No valid stock codes were resolved from the selected input source.")
    return unique, len(raw_codes) - len(unique)


def parse_execution_date(value: str | None) -> tuple[date, bool]:
    if value is None:
        return date.today(), False
    try:
        parsed = datetime.strptime(value, "%Y-%m-%d").date()
    except ValueError as exc:
        raise PreflightError("Execution date must use YYYY-MM-DD format.") from exc
    if parsed > date.today():
        raise PreflightError(f"Execution date cannot be in the future: {parsed}")
    return parsed, True


def validate_output_path(output_path: Path, input_path: Path | None, overwrite: bool) -> None:
    if output_path.suffix.lower() != ".xlsx":
        raise PreflightError(f"Output file must use the .xlsx extension: {output_path}")
    if input_path is not None and output_path == input_path:
        raise PreflightError("Input and output files must be different paths.")
    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    except Exception as exc:
        raise PreflightError(f"Unable to create output directory '{output_path.parent}': {exc}") from exc
    if output_path.exists():
        if not output_path.is_file():
            raise PreflightError(f"Output path is not a file: {output_path}")
        if not overwrite:
            raise PreflightError(f"Output already exists. Use --overwrite to replace it: {output_path}")
        try:
            with output_path.open("r+b"):
                pass
        except OSError as exc:
            raise PreflightError(f"Existing output is not writable or may be open in Excel: {output_path}: {exc}") from exc
    try:
        handle, probe_name = tempfile.mkstemp(prefix=".v19_write_probe_", dir=output_path.parent)
        os.close(handle)
        Path(probe_name).unlink()
    except Exception as exc:
        raise PreflightError(f"Output directory is not writable: {output_path.parent}: {exc}") from exc


def preflight(argv: Sequence[str] | None = None) -> tuple[RuntimeConfig, list[str]]:
    parser = build_argument_parser()
    args = parser.parse_args(argv)
    execution_date, explicit_date = parse_execution_date(args.execution_date)

    input_path: Path | None = None
    if args.input_file:
        input_path = Path(args.input_file).expanduser().resolve()
        raw_codes = read_symbols_from_file(input_path)
        input_source = str(input_path)
    else:
        raw_codes = _split_code_values(args.codes or [])
        input_source = "CLI codes"

    symbols, duplicate_count = normalize_and_deduplicate_symbols(raw_codes)
    if args.output:
        output_path = Path(args.output).expanduser().resolve()
    else:
        output_path = (SCRIPT_PATH.parent / "output" / f"V19_Momentum_Review_{execution_date.isoformat()}.xlsx").resolve()

    if not 1 <= args.workers <= 32:
        raise PreflightError("--workers must be between 1 and 32.")
    if not 1 <= args.max_attempts <= 10:
        raise PreflightError("--max-attempts must be between 1 and 10.")
    if not 1.0 <= args.request_timeout <= 120.0:
        raise PreflightError("--request-timeout must be between 1 and 120 seconds.")
    if not 0.0 <= args.request_interval <= 60.0:
        raise PreflightError("--request-interval must be between 0 and 60 seconds.")
    if not 0.0 <= args.retry_base_seconds <= 120.0:
        raise PreflightError("--retry-base-seconds must be between 0 and 120 seconds.")
    if args.history_lookback_days < 500:
        raise PreflightError("--history-lookback-days must be at least 500 for EMA200 reliability.")
    if args.progress_every < 1:
        raise PreflightError("--progress-every must be at least 1.")
    if not 0.0 <= args.stoch_zone_low < args.stoch_zone_high <= 100.0:
        raise PreflightError("Stochastic zone must satisfy 0 <= low < high <= 100.")

    validate_output_path(output_path, input_path, args.overwrite)
    config = RuntimeConfig(
        execution_date=execution_date,
        explicit_execution_date=explicit_date,
        input_source=input_source,
        raw_input_count=len(raw_codes),
        duplicate_count=duplicate_count,
        output_path=output_path,
        overwrite=args.overwrite,
        workers=args.workers,
        max_attempts=args.max_attempts,
        request_timeout=args.request_timeout,
        request_interval=args.request_interval,
        retry_base_seconds=args.retry_base_seconds,
        history_lookback_days=args.history_lookback_days,
        progress_every=args.progress_every,
        stoch_zone_low=args.stoch_zone_low,
        stoch_zone_high=args.stoch_zone_high,
    )
    return config, symbols


def yahoo_history_fetcher(symbol: str, start_date: date, end_date: date, timeout: float) -> pd.DataFrame:
    yf.config.debug.hide_exceptions = False
    ticker = yf.Ticker(symbol)
    return ticker.history(
        start=start_date.isoformat(),
        end=end_date.isoformat(),
        interval="1d",
        auto_adjust=True,
        actions=False,
        repair=True,
        keepna=False,
        timeout=timeout,
    )


def classify_provider_exception(exc: Exception) -> DataExtractionError:
    message = str(exc).strip() or exc.__class__.__name__
    lower = message.lower()
    if isinstance(exc, YFRateLimitError) or "429" in lower or "rate limit" in lower or "too many requests" in lower:
        return DataExtractionError("YF_RATE_LIMIT", message, provider_failure=True, retriable=True)
    if isinstance(exc, (YFPricesMissingError, YFTickerMissingError)):
        return DataExtractionError("NO_PRICE_DATA", message, provider_failure=False, retriable=False)
    if isinstance(exc, TimeoutError) or "timed out" in lower or "timeout" in lower or "curl: (28)" in lower:
        return DataExtractionError("YF_TIMEOUT", message, provider_failure=True, retriable=True)
    return DataExtractionError("YF_REQUEST_ERROR", message, provider_failure=True, retriable=True)


def fetch_history_with_retries(
    symbol: str,
    config: RuntimeConfig,
    fetcher: HistoryFetcher,
    guard: ProviderGuard,
    stats: RetryStatistics,
    *,
    sleep_fn: Callable[[float], None] = time.sleep,
    jitter_fn: Callable[[float, float], float] = random.uniform,
) -> tuple[pd.DataFrame, int]:
    start_date = config.execution_date - timedelta(days=config.history_lookback_days)
    end_date = config.execution_date
    last_error: DataExtractionError | None = None

    for attempt in range(1, config.max_attempts + 1):
        guard.wait_for_request_slot()
        stats.record_attempt()
        try:
            frame = fetcher(symbol, start_date, end_date, config.request_timeout)
            if frame is None or frame.empty:
                raise DataExtractionError("NO_PRICE_DATA", "Yahoo returned no daily price rows.")
            return frame, attempt
        except DataExtractionError as exc:
            current = exc
        except Exception as exc:
            current = classify_provider_exception(exc)

        last_error = current
        if current.code == "YF_RATE_LIMIT":
            stats.record_rate_limit()
            guard.pause_after_rate_limit(min(60.0, config.retry_base_seconds * (2 ** attempt)))
        if not current.retriable or attempt >= config.max_attempts:
            break
        stats.record_retry()
        delay = config.retry_base_seconds * (2 ** (attempt - 1)) + jitter_fn(0.0, max(0.1, config.retry_base_seconds))
        sleep_fn(delay)

    assert last_error is not None
    raise last_error


def normalize_completed_history(frame: pd.DataFrame, execution_date: date) -> pd.DataFrame:
    if frame is None or frame.empty:
        raise DataExtractionError("NO_PRICE_DATA", "No daily history was returned.")
    data = frame.copy()
    if isinstance(data.columns, pd.MultiIndex):
        flattened: list[str] = []
        for column in data.columns:
            parts = [str(part) for part in column if part not in (None, "")]
            required_part = next((part for part in parts if part in {"Open", "High", "Low", "Close", "Volume"}), None)
            flattened.append(required_part or parts[-1])
        data.columns = flattened

    required_columns = ["High", "Low", "Close"]
    missing = [column for column in required_columns if column not in data.columns]
    if missing:
        raise DataExtractionError("MISSING_REQUIRED_COLUMN", f"Missing daily columns: {', '.join(missing)}")

    try:
        index = pd.DatetimeIndex(pd.to_datetime(data.index))
    except Exception as exc:
        raise DataExtractionError("INVALID_DATE_INDEX", f"Unable to parse daily dates: {exc}") from exc
    if index.tz is not None:
        index = index.tz_localize(None)
    data.index = index.normalize()
    data = data[~data.index.duplicated(keep="last")].sort_index()
    data = data[data.index.date < execution_date]

    for column in required_columns + (["Volume"] if "Volume" in data.columns else []):
        data[column] = pd.to_numeric(data[column], errors="coerce")
    data = data.dropna(subset=required_columns)
    if data.empty:
        raise DataExtractionError("NO_COMPLETED_D1_DATA", f"No completed daily candle exists before D={execution_date}.")

    numeric = data[required_columns].to_numpy(dtype=float)
    if not np.isfinite(numeric).all():
        raise DataExtractionError("NON_FINITE_OHLC", "Daily OHLC contains non-finite values.")
    if (data["Close"] <= 0).any() or (data["High"] <= 0).any() or (data["Low"] <= 0).any():
        raise DataExtractionError("INVALID_OHLC_VALUES", "Daily OHLC contains zero or negative prices.")
    if (data["High"] < data["Low"]).any():
        raise DataExtractionError("INVALID_HIGH_LOW", "At least one daily High is below its Low.")
    if len(data) < MIN_REQUIRED_SESSIONS:
        raise DataExtractionError(
            "INSUFFICIENT_HISTORY",
            f"Only {len(data)} completed sessions are available; at least {MIN_REQUIRED_SESSIONS} are required.",
        )
    return data


def _require_finite(name: str, value: float) -> float:
    value = float(value)
    if not math.isfinite(value):
        raise DataExtractionError("INDICATOR_CALCULATION_ERROR", f"{name} is not finite.")
    return value


def _session_return(close: pd.Series, sessions: int) -> float:
    if len(close) < sessions + 1:
        raise DataExtractionError(
            "INSUFFICIENT_HISTORY",
            f"{sessions}-session return requires at least {sessions + 1} closes.",
        )
    start_price = _require_finite(f"Close[{sessions} sessions earlier]", close.iloc[-(sessions + 1)])
    end_price = _require_finite("D-1 Close", close.iloc[-1])
    if start_price <= 0:
        raise DataExtractionError("INVALID_RETURN_BASE", f"{sessions}-session starting close is not positive.")
    return (end_price / start_price) - 1.0


def calculate_stochastic(data: pd.DataFrame, config: RuntimeConfig) -> dict:
    low_n = data["Low"].rolling(window=STOCH_LOOKBACK, min_periods=STOCH_LOOKBACK).min()
    high_n = data["High"].rolling(window=STOCH_LOOKBACK, min_periods=STOCH_LOOKBACK).max()
    denominator = (high_n - low_n).replace(0.0, np.nan)
    raw_k = 100.0 * ((data["Close"] - low_n) / denominator)
    slow_k = raw_k.rolling(window=STOCH_K_SMOOTH, min_periods=STOCH_K_SMOOTH).mean()
    slow_d = slow_k.rolling(window=STOCH_D_SMOOTH, min_periods=STOCH_D_SMOOTH).mean()

    k1 = _require_finite("Stochastic %K D-1", slow_k.iloc[-1])
    d1 = _require_finite("Stochastic %D D-1", slow_d.iloc[-1])
    k2 = _require_finite("Stochastic %K D-2", slow_k.iloc[-2])
    d2 = _require_finite("Stochastic %D D-2", slow_d.iloc[-2])
    midpoint = (k1 + d1) / 2.0
    fresh_cross = bool(k2 <= d2 and k1 > d1)
    near_zone = bool(config.stoch_zone_low <= midpoint <= config.stoch_zone_high)
    recent_oversold = bool(float(slow_k.tail(STOCH_RECENT_LOOKBACK).min()) <= STOCH_OVERSOLD_LEVEL)

    if fresh_cross and near_zone:
        status = "BUY_ZONE_CROSSOVER_NEAR_50"
    elif fresh_cross and midpoint < config.stoch_zone_low:
        status = "BULLISH_CROSS_BELOW_MIDLINE"
    elif fresh_cross:
        status = "BULLISH_CROSS_ABOVE_MIDLINE"
    elif k1 > d1 and k1 >= 50.0:
        status = "BULLISH_ABOVE_50_NO_FRESH_CROSS"
    elif k1 > d1:
        status = "BULLISH_BELOW_50_NO_FRESH_CROSS"
    else:
        status = "NO_BULLISH_CROSS"

    return {
        "Stochastic_Calculated": True,
        "Stochastic_K_D1": k1,
        "Stochastic_D_D1": d1,
        "Stochastic_K_D2": k2,
        "Stochastic_D_D2": d2,
        "Stochastic_Crossover_Midpoint": midpoint,
        "Stochastic_Fresh_Bullish_Crossover": fresh_cross,
        "Stochastic_Near_50_Zone": near_zone,
        "Stochastic_Recent_Oversold": recent_oversold,
        "Stochastic_Status": status,
    }


def calculate_symbol_result(
    index: int,
    symbol: str,
    data: pd.DataFrame,
    config: RuntimeConfig,
    attempts: int,
) -> SymbolResult:
    close = data["Close"].astype(float)
    d1_close = _require_finite("D-1 Close", close.iloc[-1])
    d1_date = data.index[-1].date()
    d2_date = data.index[-2].date()

    ema_50_series = close.ewm(span=EMA_MEDIUM_SPAN, adjust=False, min_periods=EMA_MEDIUM_SPAN).mean()
    ema_200_series = close.ewm(span=EMA_LONG_SPAN, adjust=False, min_periods=EMA_LONG_SPAN).mean()
    ema_50 = _require_finite("EMA50 D-1", ema_50_series.iloc[-1])
    ema_200 = _require_finite("EMA200 D-1", ema_200_series.iloc[-1])

    ema_fast = close.ewm(span=MACD_FAST_SPAN, adjust=False, min_periods=MACD_FAST_SPAN).mean()
    ema_slow = close.ewm(span=MACD_SLOW_SPAN, adjust=False, min_periods=MACD_SLOW_SPAN).mean()
    macd_series = ema_fast - ema_slow
    signal_series = macd_series.ewm(
        span=MACD_SIGNAL_SPAN,
        adjust=False,
        min_periods=MACD_SIGNAL_SPAN,
    ).mean()
    macd_line = _require_finite("MACD line D-1", macd_series.iloc[-1])
    macd_signal = _require_finite("MACD signal D-1", signal_series.iloc[-1])
    macd_histogram = macd_line - macd_signal

    return_short = _session_return(close, SHORT_WINDOW_SESSIONS)
    return_medium = _session_return(close, MEDIUM_WINDOW_SESSIONS)
    return_long = _session_return(close, LONG_WINDOW_SESSIONS)

    price_above_ema50 = bool(d1_close > ema_50)
    price_above_ema200 = bool(d1_close > ema_200)
    bull_zone = bool(price_above_ema50 and price_above_ema200)
    macd_above_zero = bool(macd_line > 0.0)
    signal_above_zero = bool(macd_signal > 0.0)
    macd_above_signal = bool(macd_line > macd_signal)
    macd_buyer_zone = bool(macd_above_zero and signal_above_zero and macd_above_signal)

    returns = (return_short, return_medium, return_long)
    if all(value >= TIER_1_THRESHOLD for value in returns):
        momentum_band = "TIER_1_RETURN_BAND"
        proposed_tier = "Tier 1"
    elif all(value >= TIER_2_THRESHOLD for value in returns):
        momentum_band = "TIER_2_RETURN_BAND"
        proposed_tier = "Tier 2"
    else:
        momentum_band = "BELOW_TIER_2_RETURN_BAND"
        proposed_tier = None

    rejection_reasons: list[str] = []
    if not price_above_ema50:
        rejection_reasons.append("PRICE_NOT_ABOVE_EMA50")
    if not price_above_ema200:
        rejection_reasons.append("PRICE_NOT_ABOVE_EMA200")
    if not macd_above_zero:
        rejection_reasons.append("MACD_LINE_NOT_ABOVE_ZERO")
    if not signal_above_zero:
        rejection_reasons.append("MACD_SIGNAL_NOT_ABOVE_ZERO")
    if not macd_above_signal:
        rejection_reasons.append("MACD_LINE_NOT_ABOVE_SIGNAL")
    for label, value in (("SHORT", return_short), ("MEDIUM", return_medium), ("LONG", return_long)):
        if value < TIER_2_THRESHOLD:
            rejection_reasons.append(f"{label}_RETURN_BELOW_10_PERCENT")

    qualified = bool(bull_zone and macd_buyer_zone and proposed_tier is not None)
    if qualified and proposed_tier == "Tier 1":
        status = STATUS_TIER_1
    elif qualified:
        status = STATUS_TIER_2
    else:
        status = STATUS_REJECT

    stochastic = {
        "Stochastic_Calculated": False,
        "Stochastic_K_D1": None,
        "Stochastic_D_D1": None,
        "Stochastic_K_D2": None,
        "Stochastic_D_D2": None,
        "Stochastic_Crossover_Midpoint": None,
        "Stochastic_Fresh_Bullish_Crossover": None,
        "Stochastic_Near_50_Zone": None,
        "Stochastic_Recent_Oversold": None,
        "Stochastic_Status": "NOT_EVALUATED",
    }
    if qualified:
        stochastic = calculate_stochastic(data, config)

    tier_value = proposed_tier if qualified else None
    technical = {
        "Symbol": symbol,
        "Execution_Date_D": config.execution_date,
        "D1_Date": d1_date,
        "D2_Date": d2_date,
        "D1_Close": d1_close,
        "EMA_50": ema_50,
        "EMA_200": ema_200,
        "Price_vs_EMA50_%": (d1_close / ema_50) - 1.0,
        "Price_vs_EMA200_%": (d1_close / ema_200) - 1.0,
        "Price_Above_EMA50": price_above_ema50,
        "Price_Above_EMA200": price_above_ema200,
        "Bull_Zone_Status": bull_zone,
        "MACD_Line": macd_line,
        "MACD_Signal": macd_signal,
        "MACD_Histogram": macd_histogram,
        "MACD_Line_Above_Zero": macd_above_zero,
        "MACD_Signal_Above_Zero": signal_above_zero,
        "MACD_Line_Above_Signal": macd_above_signal,
        "MACD_Buyer_Zone": macd_buyer_zone,
        "Short_Window_Sessions": SHORT_WINDOW_SESSIONS,
        "Return_Short": return_short,
        "Medium_Window_Sessions": MEDIUM_WINDOW_SESSIONS,
        "Return_Medium": return_medium,
        "Long_Window_Sessions": LONG_WINDOW_SESSIONS,
        "Return_Long": return_long,
        "Momentum_Return_Band": momentum_band,
        "Qualification_Status": status,
        "Qualification_Tier": tier_value or "NOT_QUALIFIED",
        "Rejection_Reasons": "; ".join(rejection_reasons),
        **stochastic,
        "Data_Integrity_Flag": "VALID",
    }
    baseline = {
        "Symbol": symbol,
        "Baseline_Start_Date": data.index[0].date(),
        "Baseline_Start_Price": float(close.iloc[0]),
        "D1_Date": d1_date,
        "D1_Close": d1_close,
        "D2_Date": d2_date,
        "Total_Sessions_Available": len(data),
        "Execution_Date_D": config.execution_date,
        "Calendar_Days_D1_to_D": (config.execution_date - d1_date).days,
    }

    if qualified:
        message = (
            f"Qualified {proposed_tier}: D-1 close is above EMA50 and EMA200; MACD line and signal are above zero "
            f"with MACD above signal; all 21/63/126-session returns meet the {TIER_1_THRESHOLD if proposed_tier == 'Tier 1' else TIER_2_THRESHOLD:.0%} band. "
            f"Stochastic={stochastic['Stochastic_Status']}."
        )
    else:
        message = "Rejected: " + "; ".join(rejection_reasons)

    review = {
        "Symbol": symbol,
        "Status": status,
        "Qualification_Tier": tier_value or "NOT_QUALIFIED",
        "D1_Date": d1_date,
        "Bull_Zone_Status": bull_zone,
        "MACD_Buyer_Zone": macd_buyer_zone,
        "Momentum_Return_Band": momentum_band,
        "Stochastic_Status": stochastic["Stochastic_Status"],
        "Attempts": attempts,
        "Provider_Failure": False,
        "Error_Code": None,
        "Message": message,
    }

    scoring = None
    if qualified:
        scoring = {
            "Symbol": symbol,
            "Last Traded Price": d1_close,
            "D1_Date": d1_date,
            "Primary_Evaluation_Window": f"{MEDIUM_WINDOW_SESSIONS}-Session Rolling",
            "Return_Short": return_short,
            "Return_Medium": return_medium,
            "Return_Long": return_long,
            "Qualification_Tier": proposed_tier,
            "Bull_Zone_Status": bull_zone,
            "MACD_Buyer_Zone": macd_buyer_zone,
            "Stochastic_K_D1": stochastic["Stochastic_K_D1"],
            "Stochastic_D_D1": stochastic["Stochastic_D_D1"],
            "Stochastic_K_D2": stochastic["Stochastic_K_D2"],
            "Stochastic_D_D2": stochastic["Stochastic_D_D2"],
            "Stochastic_Fresh_Bullish_Crossover": stochastic["Stochastic_Fresh_Bullish_Crossover"],
            "Stochastic_Near_50_Zone": stochastic["Stochastic_Near_50_Zone"],
            "Stochastic_Recent_Oversold": stochastic["Stochastic_Recent_Oversold"],
            "Stochastic_Status": stochastic["Stochastic_Status"],
        }

    return SymbolResult(
        index=index,
        symbol=symbol,
        status=status,
        tier=proposed_tier if qualified else None,
        attempts=attempts,
        provider_failure=False,
        error_code=None,
        review=review,
        baseline=baseline,
        technical=technical,
        scoring=scoring,
    )


def invalid_symbol_result(
    index: int,
    symbol: str,
    error: DataExtractionError,
    attempts: int,
) -> SymbolResult:
    review = {
        "Symbol": symbol,
        "Status": STATUS_INVALID,
        "Qualification_Tier": "NOT_AVAILABLE",
        "D1_Date": None,
        "Bull_Zone_Status": None,
        "MACD_Buyer_Zone": None,
        "Momentum_Return_Band": "NOT_AVAILABLE",
        "Stochastic_Status": "NOT_EVALUATED",
        "Attempts": attempts,
        "Provider_Failure": error.provider_failure,
        "Error_Code": error.code,
        "Message": str(error),
    }
    return SymbolResult(
        index=index,
        symbol=symbol,
        status=STATUS_INVALID,
        tier=None,
        attempts=attempts,
        provider_failure=error.provider_failure,
        error_code=error.code,
        review=review,
    )


def evaluate_symbol(
    index: int,
    symbol: str,
    config: RuntimeConfig,
    fetcher: HistoryFetcher,
    guard: ProviderGuard,
    stats: RetryStatistics,
) -> SymbolResult:
    attempts = 0
    try:
        raw, attempts = fetch_history_with_retries(symbol, config, fetcher, guard, stats)
        completed = normalize_completed_history(raw, config.execution_date)
        return calculate_symbol_result(index, symbol, completed, config, attempts)
    except DataExtractionError as exc:
        return invalid_symbol_result(index, symbol, exc, attempts or config.max_attempts if exc.retriable else attempts or 1)
    except Exception as exc:
        error = DataExtractionError(
            "INTERNAL_WORKER_ERROR",
            f"Unexpected worker failure: {exc}",
            provider_failure=False,
        )
        return invalid_symbol_result(index, symbol, error, attempts or 1)


def run_scan(
    symbols: Sequence[str],
    config: RuntimeConfig,
    *,
    fetcher: HistoryFetcher = yahoo_history_fetcher,
    guard: ProviderGuard | None = None,
    stats: RetryStatistics | None = None,
) -> tuple[list[SymbolResult], RetryStatistics]:
    shared_guard = guard or ProviderGuard(config.request_interval)
    shared_stats = stats or RetryStatistics()
    results: list[SymbolResult] = []
    total = len(symbols)
    started = time.monotonic()
    tier_1_count = 0
    tier_2_count = 0
    rejected_count = 0
    invalid_count = 0

    with ThreadPoolExecutor(max_workers=config.workers, thread_name_prefix="v19-yf") as executor:
        futures = {
            executor.submit(evaluate_symbol, index, symbol, config, fetcher, shared_guard, shared_stats): symbol
            for index, symbol in enumerate(symbols)
        }
        for completed_count, future in enumerate(as_completed(futures), start=1):
            symbol = futures[future]
            try:
                result = future.result()
            except Exception as exc:
                result = invalid_symbol_result(
                    next(index for index, candidate in enumerate(symbols) if candidate == symbol),
                    symbol,
                    DataExtractionError("INTERNAL_FUTURE_ERROR", str(exc)),
                    1,
                )
            results.append(result)
            if result.status == STATUS_TIER_1:
                tier_1_count += 1
            elif result.status == STATUS_TIER_2:
                tier_2_count += 1
            elif result.status == STATUS_REJECT:
                rejected_count += 1
            else:
                invalid_count += 1
            if (
                completed_count == 1
                or completed_count == total
                or completed_count % config.progress_every == 0
                or result.status == STATUS_INVALID
            ):
                elapsed = time.monotonic() - started
                print(
                    f"[PROGRESS] {completed_count}/{total} completed | processed={tier_1_count + tier_2_count + rejected_count} "
                    f"| tier1={tier_1_count} tier2={tier_2_count} rejected={rejected_count} invalid={invalid_count} "
                    f"| latest={symbol}:{result.status} | elapsed={elapsed:.1f}s",
                    flush=True,
                )

    results.sort(key=lambda item: item.index)
    return results, shared_stats


def determine_run_status(results: Sequence[SymbolResult]) -> str:
    processed = sum(result.status != STATUS_INVALID for result in results)
    provider_failures = sum(result.status == STATUS_INVALID and result.provider_failure for result in results)
    if processed == 0:
        return RUN_INVALID
    if provider_failures > 0:
        return RUN_PARTIAL
    return RUN_COMPLETE


def build_summary_record(
    config: RuntimeConfig,
    symbols: Sequence[str],
    results: Sequence[SymbolResult],
    stats: RetryStatistics,
    started_utc: datetime,
    elapsed_seconds: float,
) -> dict:
    tier_1 = sum(result.status == STATUS_TIER_1 for result in results)
    tier_2 = sum(result.status == STATUS_TIER_2 for result in results)
    rejected = sum(result.status == STATUS_REJECT for result in results)
    invalid = sum(result.status == STATUS_INVALID for result in results)
    processed = tier_1 + tier_2 + rejected
    provider_failures = sum(result.status == STATUS_INVALID and result.provider_failure for result in results)
    d1_dates = [result.baseline["D1_Date"] for result in results if result.baseline is not None]
    run_status = determine_run_status(results)
    return {
        "engine_name": ENGINE_NAME,
        "engine_version": ENGINE_VERSION,
        "script_name": SCRIPT_PATH.name,
        "script_path": str(SCRIPT_PATH),
        "execution_started_utc": started_utc.isoformat(),
        "execution_finished_utc": datetime.now(timezone.utc).isoformat(),
        "elapsed_seconds": round(elapsed_seconds, 3),
        "run_mode": config.run_mode,
        "execution_date_D": config.execution_date,
        "candle_policy": "Use only completed daily candles dated before D",
        "data_through_earliest": min(d1_dates) if d1_dates else None,
        "data_through_latest": max(d1_dates) if d1_dates else None,
        "input_source": config.input_source,
        "total_input_read": config.raw_input_count,
        "total_unique_codes": len(symbols),
        "duplicate_codes_removed": config.duplicate_count,
        "total_processed": processed,
        "status_evaluated": tier_1 + tier_2,
        "qualified_tier_1": tier_1,
        "qualified_tier_2": tier_2,
        "status_rejected": rejected,
        "status_failed": invalid,
        "status_invalid_data": invalid,
        "provider_failure_count": provider_failures,
        "coverage_percent": (processed / len(symbols)) if symbols else 0.0,
        "request_attempts": stats.request_attempts,
        "retry_count": stats.retries,
        "rate_limit_events": stats.rate_limit_events,
        "workers": config.workers,
        "max_attempts_per_symbol": config.max_attempts,
        "request_timeout_seconds": config.request_timeout,
        "request_interval_seconds": config.request_interval,
        "short_window_sessions": SHORT_WINDOW_SESSIONS,
        "medium_window_sessions": MEDIUM_WINDOW_SESSIONS,
        "long_window_sessions": LONG_WINDOW_SESSIONS,
        "tier_1_threshold": TIER_1_THRESHOLD,
        "tier_2_threshold": TIER_2_THRESHOLD,
        "stochastic_parameters": f"{STOCH_LOOKBACK},{STOCH_K_SMOOTH},{STOCH_D_SMOOTH}",
        "stochastic_near_50_zone": f"{config.stoch_zone_low:.2f}-{config.stoch_zone_high:.2f}",
        "run_status": run_status,
        "output_path": str(config.output_path),
    }


def _frame(records: list[dict], columns: Sequence[str]) -> pd.DataFrame:
    return pd.DataFrame(records, columns=list(columns))


def _apply_workbook_style(workbook: openpyxl.Workbook) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    qualified_fill = PatternFill("solid", fgColor="E2F0D9")
    reject_fill = PatternFill("solid", fgColor="FFF2CC")
    invalid_fill = PatternFill("solid", fgColor="FCE4D6")

    percent_headers = {
        "coverage_percent",
        "tier_1_threshold",
        "tier_2_threshold",
        "Price_vs_EMA50_%",
        "Price_vs_EMA200_%",
        "Return_Short",
        "Return_Medium",
        "Return_Long",
    }
    date_headers = {
        "execution_date_D",
        "data_through_earliest",
        "data_through_latest",
        "Execution_Date_D",
        "Baseline_Start_Date",
        "D1_Date",
        "D2_Date",
    }
    price_headers = {"D1_Close", "EMA_50", "EMA_200", "Baseline_Start_Price", "Last Traded Price"}
    decimal_headers = {
        "MACD_Line",
        "MACD_Signal",
        "MACD_Histogram",
        "Stochastic_K_D1",
        "Stochastic_D_D1",
        "Stochastic_K_D2",
        "Stochastic_D_D2",
        "Stochastic_Crossover_Midpoint",
    }

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False
        worksheet.freeze_panes = "A2"
        if worksheet.max_column:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 32

        headers = [cell.value for cell in worksheet[1]]
        for column_index, header in enumerate(headers, start=1):
            if header in percent_headers:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_index).number_format = "0.00%"
            elif header in date_headers:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_index).number_format = "yyyy-mm-dd"
            elif header in price_headers:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_index).number_format = "0.0000"
            elif header in decimal_headers:
                for row in range(2, worksheet.max_row + 1):
                    worksheet.cell(row=row, column=column_index).number_format = "0.0000"

            max_length = len(str(header or ""))
            for row in worksheet.iter_rows(
                min_row=2,
                max_row=min(worksheet.max_row, 202),
                min_col=column_index,
                max_col=column_index,
            ):
                value = row[0].value
                if value is not None:
                    max_length = max(max_length, len(str(value)))
            worksheet.column_dimensions[get_column_letter(column_index)].width = min(max(max_length + 2, 10), 48)

        status_column = next((index + 1 for index, header in enumerate(headers) if header in {"Status", "Qualification_Status"}), None)
        if status_column:
            for row in range(2, worksheet.max_row + 1):
                value = str(worksheet.cell(row=row, column=status_column).value or "")
                fill = None
                if value.startswith("QUALIFIED"):
                    fill = qualified_fill
                elif value == STATUS_REJECT:
                    fill = reject_fill
                elif value == STATUS_INVALID:
                    fill = invalid_fill
                if fill:
                    worksheet.cell(row=row, column=status_column).fill = fill


def write_output_workbook(
    config: RuntimeConfig,
    symbols: Sequence[str],
    results: Sequence[SymbolResult],
    summary_record: dict,
) -> None:
    review = _frame([result.review for result in results], REVIEW_COLUMNS)
    baseline = _frame([result.baseline for result in results if result.baseline is not None], BASELINE_COLUMNS)
    technical = _frame([result.technical for result in results if result.technical is not None], TECHNICAL_COLUMNS)
    scoring = _frame([result.scoring for result in results if result.scoring is not None], SCORING_COLUMNS)
    if not scoring.empty:
        tier_order = scoring["Qualification_Tier"].map({"Tier 1": 1, "Tier 2": 2}).fillna(3)
        scoring = (
            scoring.assign(_Tier_Order=tier_order)
            .sort_values(["_Tier_Order", "Return_Medium"], ascending=[True, False], kind="stable")
            .drop(columns=["_Tier_Order"])
        )

    output_path = config.output_path
    descriptor, temporary_name = tempfile.mkstemp(prefix=".v19_output_", suffix=".xlsx", dir=output_path.parent)
    os.close(descriptor)
    temporary_path = Path(temporary_name)
    try:
        with pd.ExcelWriter(temporary_path, engine="openpyxl") as writer:
            pd.DataFrame([summary_record]).to_excel(writer, sheet_name="Summary", index=False)
            review.to_excel(writer, sheet_name="Review", index=False)
            baseline.to_excel(writer, sheet_name="Baseline Info", index=False)
            technical.to_excel(writer, sheet_name="Technical Data", index=False)
            scoring.to_excel(writer, sheet_name="Scoring", index=False)

        workbook = openpyxl.load_workbook(temporary_path)
        _apply_workbook_style(workbook)
        workbook.save(temporary_path)
        workbook.close()

        verification = openpyxl.load_workbook(temporary_path, read_only=True, data_only=False)
        if verification.sheetnames != SHEET_NAMES:
            found = ", ".join(verification.sheetnames)
            verification.close()
            raise RuntimeError(f"Workbook schema validation failed; found sheets: {found}")
        verification.close()
        os.replace(temporary_path, output_path)
    except Exception:
        if temporary_path.exists():
            temporary_path.unlink()
        raise


def print_initial_baseline(config: RuntimeConfig, symbols: Sequence[str]) -> None:
    print("=" * 88)
    print(f"{ENGINE_NAME} ({ENGINE_VERSION})")
    print("=" * 88)
    print(f"Script Name                 : {SCRIPT_PATH.name}")
    print(f"Script Location             : {SCRIPT_PATH}")
    print(f"Run Mode                    : {config.run_mode}")
    print(f"Execution Date D            : {config.execution_date}")
    print("Daily Candle Baseline       : latest completed session strictly before D")
    print(f"Input Source                : {config.input_source}")
    print(f"Codes Read                  : {config.raw_input_count}")
    print(f"Unique Codes to Process     : {len(symbols)}")
    print(f"Duplicates Removed          : {config.duplicate_count}")
    print(f"Workers                     : {config.workers}")
    print(f"Max Attempts per Code       : {config.max_attempts}")
    print(f"Request Timeout             : {config.request_timeout:.1f} seconds")
    print(f"Minimum Request Interval    : {config.request_interval:.2f} seconds")
    print(f"Output File                 : {config.output_path}")
    print("Starting Counts             : processed=0 tier1=0 tier2=0 rejected=0 failed=0")
    print("Point 1                     : D-1 Close > EMA50 and EMA200")
    print("Point 2                     : MACD > 0, Signal > 0, MACD > Signal")
    print("Point 3                     : all 21/63/126 returns >=20% (Tier 1) or >=10% (Tier 2)")
    print("Point 4                     : stochastic 14,3,3 calculated for qualified codes only")
    print("-" * 88, flush=True)


def print_final_summary(summary: dict) -> None:
    print("-" * 88)
    print("V19 EXECUTION SUMMARY")
    print("-" * 88)
    print(f"Run Status                  : {summary['run_status']}")
    print(f"Codes Read / Unique         : {summary['total_input_read']} / {summary['total_unique_codes']}")
    print(f"Processed                   : {summary['total_processed']}")
    print(f"Qualified Tier 1            : {summary['qualified_tier_1']}")
    print(f"Qualified Tier 2            : {summary['qualified_tier_2']}")
    print(f"Rejected                    : {summary['status_rejected']}")
    print(f"Failed / Invalid Data       : {summary['status_invalid_data']}")
    print(f"Provider Failures           : {summary['provider_failure_count']}")
    print(f"Requests / Retries / 429s   : {summary['request_attempts']} / {summary['retry_count']} / {summary['rate_limit_events']}")
    print(f"Data Through                : {summary['data_through_earliest']} to {summary['data_through_latest']}")
    print(f"Output File                 : {summary['output_path']}")
    print("=" * 88, flush=True)


def temporary_yfinance_cache():
    """Return a TemporaryDirectory and point yfinance technical cache at it."""
    directory = tempfile.TemporaryDirectory(prefix="v19_yfinance_cache_", ignore_cleanup_errors=True)
    os.environ["YFINANCE_CACHE_DIR"] = directory.name
    try:
        yf.set_tz_cache_location(directory.name)
    except Exception:
        pass
    return directory


def close_yfinance_cache_handles() -> None:
    """Close Yahoo SQLite handles so Windows can remove the per-run cache."""
    try:
        import yfinance.cache as yf_cache

        for getter in (yf_cache.get_cookie_cache, yf_cache.get_tz_cache, yf_cache.get_isin_cache):
            cache = getter()
            database = getattr(cache, "db", None)
            if database is None:
                continue
            try:
                if not database.is_closed():
                    database.close()
            except Exception:
                pass
    except Exception:
        pass


def cleanup_yfinance_cache(directory) -> None:
    """Best-effort cleanup must never prevent the workbook from being written."""
    close_yfinance_cache_handles()
    try:
        directory.cleanup()
    except Exception as exc:
        print(f"[CACHE WARNING] Temporary Yahoo cache cleanup was incomplete: {exc}", file=sys.stderr)
    finally:
        os.environ.pop("YFINANCE_CACHE_DIR", None)


def main(argv: Sequence[str] | None = None) -> int:
    try:
        config, symbols = preflight(argv)
    except PreflightError as exc:
        print(f"[PRE-FLIGHT ERROR] {exc}", file=sys.stderr)
        return 2

    print_initial_baseline(config, symbols)
    started_utc = datetime.now(timezone.utc)
    started_monotonic = time.monotonic()
    cache_directory = temporary_yfinance_cache()
    try:
        results, stats = run_scan(symbols, config)
    finally:
        cleanup_yfinance_cache(cache_directory)

    elapsed_seconds = time.monotonic() - started_monotonic
    summary = build_summary_record(config, symbols, results, stats, started_utc, elapsed_seconds)
    try:
        write_output_workbook(config, symbols, results, summary)
    except Exception as exc:
        print(f"[OUTPUT ERROR] Unable to write '{config.output_path}': {exc}", file=sys.stderr)
        return 5

    print_final_summary(summary)
    if summary["run_status"] == RUN_INVALID:
        return 4
    if summary["run_status"] == RUN_PARTIAL:
        return 3
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
