import sys
import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


SCANNER_DIR = Path(__file__).resolve().parents[1]
if str(SCANNER_DIR) not in sys.path:
    sys.path.insert(0, str(SCANNER_DIR))

import Live_Scanner_v19 as v19


def synthetic_history(execution_date: date, daily_growth: float, sessions: int = 260) -> pd.DataFrame:
    all_dates = pd.bdate_range(end=pd.Timestamp(execution_date), periods=sessions + 1)
    closes = 20.0 * np.power(1.0 + daily_growth, np.arange(len(all_dates), dtype=float))
    return pd.DataFrame(
        {
            "Open": closes * 0.997,
            "High": closes * 1.02,
            "Low": closes * 0.98,
            "Close": closes,
            "Volume": np.full(len(closes), 1_000_000.0),
        },
        index=all_dates,
    )


def runtime_config(output_path: Path, execution_date: date = date(2026, 8, 6), **overrides) -> v19.RuntimeConfig:
    values = {
        "execution_date": execution_date,
        "explicit_execution_date": True,
        "input_source": "unit test",
        "raw_input_count": 1,
        "duplicate_count": 0,
        "output_path": output_path,
        "overwrite": False,
        "workers": 2,
        "max_attempts": 2,
        "request_timeout": 5.0,
        "request_interval": 0.0,
        "retry_base_seconds": 0.0,
        "history_lookback_days": 1_100,
        "progress_every": 100,
        "stoch_zone_low": 45.0,
        "stoch_zone_high": 55.0,
    }
    values.update(overrides)
    return v19.RuntimeConfig(**values)


class V19CompletedCandleTests(unittest.TestCase):
    def test_session_return_uses_the_full_number_of_intervals(self):
        closes = pd.Series(np.arange(1.0, 23.0))
        calculated = v19._session_return(closes, 21)
        self.assertAlmostEqual(calculated, 22.0 / 1.0 - 1.0)

    def test_execution_date_candle_is_excluded(self):
        execution_date = date(2026, 8, 6)
        frame = synthetic_history(execution_date, 0.005)
        self.assertIn(pd.Timestamp(execution_date), frame.index)

        normalized = v19.normalize_completed_history(frame, execution_date)

        self.assertLess(normalized.index[-1].date(), execution_date)
        self.assertEqual(normalized.index[-1].date(), date(2026, 8, 5))

    def test_insufficient_history_is_invalid(self):
        frame = synthetic_history(date(2026, 8, 6), 0.005, sessions=100)
        with self.assertRaises(v19.DataExtractionError) as captured:
            v19.normalize_completed_history(frame, date(2026, 8, 6))
        self.assertEqual(captured.exception.code, "INSUFFICIENT_HISTORY")


class V19QualificationTests(unittest.TestCase):
    def calculate(self, daily_growth: float):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            config = runtime_config(Path(directory) / "output.xlsx", execution_date)
            data = v19.normalize_completed_history(synthetic_history(execution_date, daily_growth), execution_date)
            return v19.calculate_symbol_result(0, "TEST", data, config, 1)

    def test_tier_one_requires_all_three_returns_at_least_twenty_percent(self):
        result = self.calculate(0.01)
        self.assertEqual(result.status, v19.STATUS_TIER_1)
        self.assertEqual(result.tier, "Tier 1")
        self.assertTrue(result.technical["Bull_Zone_Status"])
        self.assertTrue(result.technical["MACD_Buyer_Zone"])
        self.assertGreaterEqual(result.technical["Return_Short"], 0.20)
        self.assertGreaterEqual(result.technical["Return_Medium"], 0.20)
        self.assertGreaterEqual(result.technical["Return_Long"], 0.20)
        self.assertTrue(result.technical["Stochastic_Calculated"])
        self.assertNotIn("EMA_20", result.technical)

    def test_tier_two_requires_all_three_returns_at_least_ten_percent(self):
        result = self.calculate(0.005)
        self.assertEqual(result.status, v19.STATUS_TIER_2)
        self.assertEqual(result.tier, "Tier 2")
        self.assertGreaterEqual(result.technical["Return_Short"], 0.10)
        self.assertLess(result.technical["Return_Short"], 0.20)
        self.assertTrue(result.technical["Stochastic_Calculated"])

    def test_below_tier_two_is_rejected_and_stochastic_is_skipped(self):
        result = self.calculate(0.003)
        self.assertEqual(result.status, v19.STATUS_REJECT)
        self.assertIsNone(result.tier)
        self.assertEqual(result.technical["Momentum_Return_Band"], "BELOW_TIER_2_RETURN_BAND")
        self.assertFalse(result.technical["Stochastic_Calculated"])
        self.assertEqual(result.technical["Stochastic_Status"], "NOT_EVALUATED")
        self.assertIn("SHORT_RETURN_BELOW_10_PERCENT", result.technical["Rejection_Reasons"])

    def test_negative_trend_records_all_failed_core_conditions(self):
        result = self.calculate(-0.003)
        self.assertEqual(result.status, v19.STATUS_REJECT)
        self.assertFalse(result.technical["Bull_Zone_Status"])
        self.assertFalse(result.technical["MACD_Buyer_Zone"])
        reasons = result.technical["Rejection_Reasons"]
        self.assertIn("PRICE_NOT_ABOVE_EMA50", reasons)
        self.assertIn("PRICE_NOT_ABOVE_EMA200", reasons)
        self.assertIn("MACD_LINE_NOT_ABOVE_ZERO", reasons)


class V19InputAndRetryTests(unittest.TestCase):
    def test_csv_file_input_reads_first_column_and_removes_header(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "codes.csv"
            output_path = Path(directory) / "review.xlsx"
            input_path.write_text("Symbol,Note\nAAPL,one\nMSFT,two\nAAPL,duplicate\n", encoding="utf-8")

            config, symbols = v19.preflight(["-i", str(input_path), "-o", str(output_path)])

        self.assertEqual(symbols, ["AAPL", "MSFT"])
        self.assertEqual(config.raw_input_count, 3)
        self.assertEqual(config.duplicate_count, 1)

    def test_xlsx_file_input_reads_first_sheet_first_column(self):
        with tempfile.TemporaryDirectory() as directory:
            input_path = Path(directory) / "codes.xlsx"
            output_path = Path(directory) / "review.xlsx"
            frame = pd.DataFrame({"Ticker": ["AAPL", "MSFT"], "Note": ["one", "two"]})
            frame.to_excel(input_path, index=False)

            _config, symbols = v19.preflight(["-i", str(input_path), "-o", str(output_path)])

        self.assertEqual(symbols, ["AAPL", "MSFT"])

    def test_preflight_deduplicates_codes_and_validates_output_before_run(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "review.xlsx"
            config, symbols = v19.preflight(
                [
                    "-c",
                    "AAPL,MSFT",
                    "AAPL",
                    "-d",
                    "2026-08-05",
                    "-o",
                    str(output),
                    "--workers",
                    "2",
                ]
            )
        self.assertEqual(symbols, ["AAPL", "MSFT"])
        self.assertEqual(config.raw_input_count, 3)
        self.assertEqual(config.duplicate_count, 1)
        self.assertEqual(config.run_mode, "BACKTEST_D1")

    def test_existing_output_requires_explicit_overwrite(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "review.xlsx"
            output.write_bytes(b"existing")
            with self.assertRaises(v19.PreflightError):
                v19.preflight(["-c", "AAPL", "-o", str(output)])

    def test_timeout_is_retried_then_succeeds(self):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            config = runtime_config(Path(directory) / "output.xlsx", execution_date)
            calls = []

            def flaky_fetcher(symbol, start_date, end_date, timeout):
                calls.append(symbol)
                if len(calls) == 1:
                    raise TimeoutError("simulated timeout")
                return synthetic_history(execution_date, 0.005)

            stats = v19.RetryStatistics()
            guard = v19.ProviderGuard(0.0, sleep_fn=lambda _: None)
            frame, attempts = v19.fetch_history_with_retries(
                "TEST",
                config,
                flaky_fetcher,
                guard,
                stats,
                sleep_fn=lambda _: None,
                jitter_fn=lambda _a, _b: 0.0,
            )

        self.assertFalse(frame.empty)
        self.assertEqual(attempts, 2)
        self.assertEqual(stats.request_attempts, 2)
        self.assertEqual(stats.retries, 1)


class V19ConcurrencyAndWorkbookTests(unittest.TestCase):
    def test_thread_completion_does_not_change_input_order(self):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            config = runtime_config(Path(directory) / "output.xlsx", execution_date, raw_input_count=3)

            def fetcher(symbol, start_date, end_date, timeout):
                growth = {"FIRST": 0.005, "SECOND": 0.01, "THIRD": 0.003}[symbol]
                return synthetic_history(execution_date, growth)

            results, _ = v19.run_scan(["FIRST", "SECOND", "THIRD"], config, fetcher=fetcher)

        self.assertEqual([result.symbol for result in results], ["FIRST", "SECOND", "THIRD"])

    def test_invalid_symbol_does_not_stop_remaining_symbols(self):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            config = runtime_config(Path(directory) / "output.xlsx", execution_date, raw_input_count=2)

            def fetcher(symbol, start_date, end_date, timeout):
                if symbol == "BROKEN":
                    raise v19.DataExtractionError("NO_PRICE_DATA", "simulated missing data")
                return synthetic_history(execution_date, 0.005)

            results, _ = v19.run_scan(["BROKEN", "GOOD"], config, fetcher=fetcher)

        self.assertEqual(results[0].status, v19.STATUS_INVALID)
        self.assertEqual(results[0].error_code, "NO_PRICE_DATA")
        self.assertEqual(results[1].status, v19.STATUS_TIER_2)
        self.assertEqual(v19.determine_run_status(results), v19.RUN_COMPLETE)

    def test_provider_failure_marks_partial_run(self):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            config = runtime_config(
                Path(directory) / "output.xlsx",
                execution_date,
                raw_input_count=2,
                max_attempts=1,
            )

            def fetcher(symbol, start_date, end_date, timeout):
                if symbol == "BLOCKED":
                    raise TimeoutError("simulated Yahoo timeout")
                return synthetic_history(execution_date, 0.005)

            results, _ = v19.run_scan(["BLOCKED", "GOOD"], config, fetcher=fetcher)

        self.assertTrue(results[0].provider_failure)
        self.assertEqual(v19.determine_run_status(results), v19.RUN_PARTIAL)

    def test_workbook_uses_five_sheet_base_contract(self):
        execution_date = date(2026, 8, 6)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "review.xlsx"
            config = runtime_config(output, execution_date, raw_input_count=2)
            tier_one_data = v19.normalize_completed_history(synthetic_history(execution_date, 0.01), execution_date)
            reject_data = v19.normalize_completed_history(synthetic_history(execution_date, 0.003), execution_date)
            results = [
                v19.calculate_symbol_result(0, "TIERONE", tier_one_data, config, 1),
                v19.calculate_symbol_result(1, "REJECTED", reject_data, config, 1),
            ]
            stats = v19.RetryStatistics()
            summary = v19.build_summary_record(
                config,
                ["TIERONE", "REJECTED"],
                results,
                stats,
                datetime.now(timezone.utc),
                1.0,
            )
            v19.write_output_workbook(config, ["TIERONE", "REJECTED"], results, summary)

            workbook = openpyxl.load_workbook(output, read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, v19.SHEET_NAMES)
            self.assertEqual(workbook["Review"].max_row, 3)
            self.assertEqual(workbook["Technical Data"].max_row, 3)
            self.assertEqual(workbook["Scoring"].max_row, 2)
            scoring_headers = [cell.value for cell in workbook["Scoring"][1]]
            self.assertEqual(scoring_headers, v19.SCORING_COLUMNS)
            workbook.close()


if __name__ == "__main__":
    unittest.main()
