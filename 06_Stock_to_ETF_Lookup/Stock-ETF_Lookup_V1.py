import argparse
import os
import time
import datetime
import requests
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from concurrent.futures import ThreadPoolExecutor, as_completed

# Global configuration paths and fallback defaults
DEFAULT_ETF_CSV = "D:\\Tools\\00_StockCodeMaster\\03_ETF\\22-07-US_ETF_Master_Library.csv"
FAIL_SAFE_STOCKCODE = "MU"

# Dynamic date string resolution for output tracking (ddmm format)
current_date_str = datetime.datetime.now().strftime("%d%m")
DEFAULT_OUTPUT_XLSX = f"StockCode_ETF_Lookup-{current_date_str}.xlsx"

# Standard headers matching current browser definitions
STANDARD_HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,image/apng,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.9',
    'Connection': 'keep-alive'
}

class YahooSessionManager:
    """Manages session initialization including crumb and cookie tokens to bypass blocks cleanly."""
    def __init__(self):
        self.session = requests.Session()
        self.session.headers.update(STANDARD_HEADERS)
        self.crumb = None
        self._initialize_session()

    def _initialize_session(self):
        try:
            response = self.session.get("https://finance.yahoo.com", timeout=10)
            if response.status_code == 200:
                crumb_url = "https://query2.finance.yahoo.com/v1/test/getcrumb"
                crumb_res = self.session.get(crumb_url, timeout=5)
                if crumb_res.status_code == 200 and crumb_res.text:
                    self.crumb = crumb_res.text.strip()
        except Exception:
            pass

# Initialize global tracking session framework
YF_SESSION = YahooSessionManager()

def fetch_etf_holdings_direct_api(etf_ticker):
    """
    Directly queries the public Yahoo Finance quoteSummary API module 
    to fetch constituent holdings data using session properties.
    """
    crumb_str = f"&crumb={YF_SESSION.crumb}" if YF_SESSION.crumb else ""
    url = f"https://query2.finance.yahoo.com/v10/finance/quoteSummary/{etf_ticker}?modules=topHoldings,price,institutionOwnership{crumb_str}"
    try:
        response = YF_SESSION.session.get(url, timeout=7)
        if response.status_code == 200:
            json_data = response.json()
            result = json_data.get('quoteSummary', {}).get('result', [])
            if result:
                summary = result[0]
                top_holdings = summary.get('topHoldings', {}).get('holdings', [])
                price_module = summary.get('price', {})
                
                if not top_holdings:
                    inst_ownership = summary.get('institutionOwnership', {}).get('ownershipList', [])
                    for item in inst_ownership:
                        top_holdings.append({
                            'symbol': item.get('organization'),
                            'percent': item.get('pctHeld')
                        })
                
                ltp = price_module.get('regularMarketPrice', {}).get('raw') or price_module.get('regularMarketPreviousClose', {}).get('raw') or 0.0
                currency = price_module.get('currency', 'USD')
                
                holdings_dict = {}
                for holding in top_holdings:
                    symbol = str(holding.get('symbol', '')).strip().upper()
                    
                    # Fixed structural parsing for weights across variant formats
                    percent_val = holding.get('percent', 0.0)
                    if isinstance(percent_val, dict):
                        weight = float(percent_val.get('raw', 0.0))
                    else:
                        weight = float(percent_val) if percent_val else 0.0
                        
                    if symbol:
                        if 0.0 < weight < 1.0:
                            weight *= 100
                        holdings_dict[symbol] = weight
                        
                return {'ltp': ltp, 'currency': currency, 'holdings': holdings_dict}
    except Exception:
        pass
    return None

def fetch_stock_metadata_direct(stock_ticker):
    """Queries public endpoints for stock metrics using a highly resilient dual-endpoint framework."""
    crumb_str = f"&crumb={YF_SESSION.crumb}" if YF_SESSION.crumb else ""
    url_primary = f"https://query2.finance.yahoo.com/v10/finance/quoteSummary/{stock_ticker}?modules=price,summaryDetail{crumb_str}"
    try:
        response = YF_SESSION.session.get(url_primary, timeout=7)
        if response.status_code == 200:
            json_data = response.json()
            result = json_data.get('quoteSummary', {}).get('result', [])
            if result:
                summary = result[0]
                price_mod = summary.get('price', {})
                detail_mod = summary.get('summaryDetail', {})
                
                name = price_mod.get('longName') or price_mod.get('shortName') or 'Unknown Security'
                exchange = price_mod.get('exchangeName') or 'Unknown'
                currency = price_mod.get('currency') or 'USD'
                ltp = price_mod.get('regularMarketPrice', {}).get('raw') or price_mod.get('regularMarketPreviousClose', {}).get('raw') or 0.0
                
                market_cap_raw = detail_mod.get('marketCap', {}).get('raw') or price_mod.get('marketCap', {}).get('raw') or 0
                currency_symbol = "₹" if currency == "INR" or ".NS" in stock_ticker else "$"
                
                if market_cap_raw >= 1e12:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e12:.2f}T"
                elif market_cap_raw >= 1e9:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e9:.2f}B"
                else:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e6:.2f}M" if market_cap_raw else "-"
                    
                if ltp > 0.0 and name != 'Unknown Security':
                    return {'name': name, 'exchange': exchange, 'market_cap': market_cap, 'ltp': ltp, 'currency': currency}
    except Exception:
        pass

    url_secondary = f"https://query2.finance.yahoo.com/v7/finance/quote?symbols={stock_ticker}{crumb_str}"
    try:
        response = YF_SESSION.session.get(url_secondary, timeout=7)
        if response.status_code == 200:
            json_data = response.json()
            result = json_data.get('quoteResponse', {}).get('result', [])
            if result:
                stock_info = result[0]
                name = stock_info.get('longName') or stock_info.get('shortName') or 'Unknown Security'
                exchange = stock_info.get('fullExchangeName') or stock_info.get('exchange', 'Unknown')
                currency = stock_info.get('currency', 'USD')
                ltp = stock_info.get('regularMarketPrice') or stock_info.get('regularMarketPreviousClose') or 0.0
                
                market_cap_raw = stock_info.get('marketCap', 0)
                currency_symbol = "₹" if currency == "INR" or ".NS" in stock_ticker else "$"
                
                if market_cap_raw >= 1e12:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e12:.2f}T"
                elif market_cap_raw >= 1e9:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e9:.2f}B"
                else:
                    market_cap = f"{currency_symbol}{market_cap_raw / 1e6:.2f}M" if market_cap_raw else "-"
                    
                return {'name': name, 'exchange': exchange, 'market_cap': market_cap, 'ltp': ltp, 'currency': currency}
    except Exception:
        pass

    return {'name': 'Unknown Security', 'exchange': 'Unknown', 'market_cap': '-', 'ltp': 0.0, 'currency': 'USD'}

def process_single_etf(yf_etf, orig_etf, target_stocks_yf, stock_mapping):
    data = fetch_etf_holdings_direct_api(yf_etf)
    if not data:
        return None
        
    holdings = data['holdings']
    matches_found = {}
    has_match = False
    
    for yf_stock in target_stocks_yf:
        orig_stock = stock_mapping[yf_stock]
        stock_clean = yf_stock.split('.')[0]
        
        weight_val = holdings.get(yf_stock) or holdings.get(stock_clean)
        if weight_val is not None:
            matches_found[orig_stock] = weight_val
            has_match = True
            
    if has_match:
        return {
            'orig_etf': orig_etf,
            'ltp': data['ltp'],
            'currency': data['currency'],
            'matches': matches_found
        }
    return None

def run_etf_stock_lookup(etf_csv_path, stock_codes_str, output_xlsx_path):
    start_time = time.time()
    
    # Strict input/output validation before processing threads begin
    print(f"[+] Running strict validation on file accessibility...")
    abs_etf_path = os.path.abspath(etf_csv_path)
    if not os.path.exists(abs_etf_path):
        print(f"[-] Error: Input ETF CSV does not exist at: {abs_etf_path}")
        return
    try:
        with open(abs_etf_path, 'r') as f:
            f.read(10)
    except IOError as e:
        print(f"[-] Error: Target ETF input CSV is locked or unreadable: {e}")
        return
        
    abs_out_path = os.path.abspath(output_xlsx_path)
    out_dir = os.path.dirname(abs_out_path)
    if not os.path.exists(out_dir):
        try:
            os.makedirs(out_dir, exist_ok=True)
        except IOError as e:
            print(f"[-] Error: Cannot generate output directory path: {e}")
            return
            
    try:
        with open(abs_out_path, 'a+') as f:
            pass
    except IOError as e:
        print(f"[-] Error: Output file path lacks read/write permissions: {e}")
        return
    print("[+] File I/O system validations completed cleanly.")
    
    # 1. Parse and sanitize target input stocks
    raw_stocks = [s.strip().upper() for s in stock_codes_str.split(",") if s.strip()]
    if not raw_stocks:
        print("[-] Error: No valid stock codes provided as inputs.")
        return
        
    target_stocks_yf = [s.replace("XNSE", ".NS") if "XNSE" in s else s for s in raw_stocks]
    stock_mapping = dict(zip(target_stocks_yf, raw_stocks)) 
    
    # 2. Parse ETF input CSV list
    try:
        etf_df = pd.read_csv(abs_etf_path)
        etf_col = [col for col in etf_df.columns if col.upper() in ['TICKER', 'ETF', 'SYMBOL', 'CODE']]
        if etf_col:
            raw_library_etfs = etf_df[etf_col[0]].dropna().astype(str).str.strip().str.upper().tolist()
        else:
            raw_library_etfs = etf_df.iloc[:, 0].dropna().astype(str).str.strip().str.upper().tolist()
    except Exception as e:
        print(f"[-] Error parsing ETF CSV file: {e}")
        return

    etf_list_yf = sorted(list(set([e.replace("XNSE", ".NS") if "XNSE" in e else e for e in raw_library_etfs])))
    etf_mapping = dict(zip(etf_list_yf, raw_library_etfs))
    
    print(f"[+] Loaded {len(etf_list_yf)} ETF codes from target CSV reference.")
    print(f"[+] Extracting profile metadata details for stock symbols...")

    stock_metadata = {}
    for yf_stock in target_stocks_yf:
        orig_stock = stock_mapping[yf_stock]
        stock_metadata[orig_stock] = fetch_stock_metadata_direct(yf_stock)
        print(f"    -> {orig_stock}: {stock_metadata[orig_stock]['name']} | Price: {stock_metadata[orig_stock]['ltp']} | Market Cap: {stock_metadata[orig_stock]['market_cap']}")

    matrix_data = {s: {} for s in raw_stocks}
    matched_etfs_set = set()
    etf_prices = {}
    etf_currencies = {}
    etf_read_count = 0
    
    MAX_WORKERS = 12
    print(f"[+] Spawning {MAX_WORKERS} concurrent worker threads across the target matrix...")
    
    with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
        future_to_etf = {
            executor.submit(
                process_single_etf, yf_etf, etf_mapping[yf_etf], target_stocks_yf, stock_mapping
            ): yf_etf for yf_etf in etf_list_yf
        }
        
        for future in as_completed(future_to_etf):
            etf_read_count += 1
            if len(etf_list_yf) > 100 and etf_read_count % 100 == 0:
                print(f"    -> Progress: {etf_read_count}/{len(etf_list_yf)} ETFs analyzed...")
                
            result = future.result()
            if result:
                orig_etf = result['orig_etf']
                etf_prices[orig_etf] = result['ltp']
                etf_currencies[orig_etf] = result['currency']
                matched_etfs_set.add(orig_etf)
                
                for orig_stock, weight in result['matches'].items():
                    matrix_data[orig_stock][orig_etf] = weight

    sorted_matched_etfs = sorted(list(matched_etfs_set))
    matched_stocks = [s for s in raw_stocks if len(matrix_data[s]) > 0]

    # 3. Generate Visual Excel Matrix Sheet Layout
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ETF Lookup Analysis"
    ws.views.sheetView[0].showGridLines = True
    
    navy_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    light_tint_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    alt_zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    
    font_white_bold = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Segoe UI", size=11)
    font_bold = Font(name="Segoe UI", size=11, bold=True)
    thin_border = Border(left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'), top=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0'))
    
    row_1_headers = ["Company Name", "Stock Ticker", "Exchange", "Market Cap", "Price"] + sorted_matched_etfs + ["ETF_Count"]
    ws.append(row_1_headers)
    
    row_2_headers = ["", "", "", "", ""]
    for etf in sorted_matched_etfs:
        row_2_headers.append(etf_prices.get(etf, 0.0))
    row_2_headers.append("")
    ws.append(row_2_headers)
    
    for r in [1, 2]:
        for col_idx in range(1, len(row_1_headers) + 1):
            cell = ws.cell(row=r, column=col_idx)
            cell.fill = navy_header_fill
            cell.font = font_white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            if r == 2 and 6 <= col_idx < 6 + len(sorted_matched_etfs):
                etf_name = sorted_matched_etfs[col_idx - 6]
                cell.number_format = '"₹"#,##0.00' if etf_currencies.get(etf_name) == 'INR' else '"$"#,##0.00'
                
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 24
    
    row_idx = 3
    for stock in raw_stocks:
        row_values = [stock_metadata[stock]['name'], stock, stock_metadata[stock]['exchange'], stock_metadata[stock]['market_cap'], stock_metadata[stock]['ltp']]
        
        valid_etf_count = 0
        for etf in sorted_matched_etfs:
            weight = matrix_data[stock].get(etf, "")
            row_values.append(weight)
            if weight != "":
                valid_etf_count += 1
            
        row_values.append(valid_etf_count)
        ws.append(row_values)
        
        is_even = (row_idx % 2 == 0)
        for col_idx in range(1, len(row_values) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_regular
            cell.border = thin_border
            
            if is_even and col_idx <= 5:
                cell.fill = alt_zebra_fill
            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx in [2, 4]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                cell.number_format = '"₹"#,##0.00' if stock_metadata[stock]['currency'] == 'INR' else '"$"#,##0.00'
            elif col_idx == len(row_values):
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_bold
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if cell.value != "":
                    cell.number_format = '0.00"%"'
                    cell.fill = light_tint_fill
                    cell.font = font_bold
        ws.row_dimensions[row_idx].height = 22
        row_idx += 1
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)
        
    wb.save(abs_out_path)
    
    elapsed_time = time.time() - start_time
    time_str = str(datetime.timedelta(seconds=int(elapsed_time)))
    
    print("\n" + "="*50)
    print("                 ETF LOOKUP REPORT EXECUTION")
    print("="*50)
    print(f"[>] Total Reference ETFs processed        : {etf_read_count}")
    print(f"[>] Total Stock codes processed           : {len(raw_stocks)}")
    print(f"[>] Total ETF and Stock matches found     : {len(sorted_matched_etfs)} ETFs generated dynamically.")
    print(f"[>] Total execution elapsed time          : {time_str}")
    print(f"[>] Output Saved Location                 : {abs_out_path}")
    print("="*50 + "\n")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Map Custom Lookups Between Stock Universes And ETF Constituents Matrices.")
    parser.add_argument('--csv', type=str, default=DEFAULT_ETF_CSV, help="Path to input ETF CSV reference")
    parser.add_argument('--stocks', type=str, default=None, help="Comma separated stock symbols to parse lookups for")
    parser.add_argument('--output', type=str, default=DEFAULT_OUTPUT_XLSX, help="Output destination spreadsheet name")
    
    args = parser.parse_args()
    resolved_stocks = args.stocks if args.stocks is not None else FAIL_SAFE_STOCKCODE
    run_etf_stock_lookup(args.csv, resolved_stocks, args.output)