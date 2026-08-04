import argparse
import os
import time
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import urllib.request
import urllib.parse
import json

# Global configuration paths and defaults matching exact requirements
DEFAULT_ETF_HOLDINGS_CSV = "D:\\Tools\\00_StockCodeMaster\\03_ETF\\22-07-ETF_Holdings_Detail.csv"
FAIL_SAFE_STOCKCODE = "RCAT"

current_date_str = datetime.datetime.now().strftime("%d%m")
DEFAULT_OUTPUT_XLSX = f"StockCode_ETF_Holdings_Lookup-{current_date_str}.xlsx"

STANDARD_HEADERS = {
    'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/125.0.0.0 Safari/537.36',
    'Accept': 'text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8',
    'Accept-Language': 'en-US,en;q=0.5',
    'Connection': 'keep-alive'
}

class YahooSessionManager:
    """Manages session cookies and crumb acquisition to ensure Yahoo Finance API requests succeed."""
    def __init__(self):
        self.cookie = None
        self.crumb = None
        self._refresh_session()

    def _refresh_session(self):
        try:
            req = urllib.request.Request("https://finance.yahoo.com", headers=STANDARD_HEADERS)
            with urllib.request.urlopen(req, timeout=8) as resp:
                set_cookie = resp.headers.get('Set-Cookie')
                if set_cookie:
                    self.cookie = set_cookie.split(';')[0]
            
            headers_with_cookie = STANDARD_HEADERS.copy()
            if self.cookie:
                headers_with_cookie['Cookie'] = self.cookie
                
            req_crumb = urllib.request.Request("https://query2.finance.yahoo.com/v1/test/getcrumb", headers=headers_with_cookie)
            with urllib.request.urlopen(req_crumb, timeout=8) as resp:
                crumb_val = resp.read().decode('utf-8').strip()
                if crumb_val and "<html>" not in crumb_val:
                    self.crumb = crumb_val
        except Exception:
            pass

YF_SESSION = YahooSessionManager()

def fetch_quotes_batch(tickers):
    """Fetches live market prices and metadata using Yahoo Finance v7 quote endpoint with robust session cookie and crumb authentication."""
    results_map = {}
    if not tickers:
        return results_map
        
    if not YF_SESSION.crumb:
        YF_SESSION._refresh_session()

    chunk_size = 50
    for i in range(0, len(tickers), chunk_size):
        chunk = tickers[i:i + chunk_size]
        symbols_str = ",".join(chunk)
        
        crumb_param = f"&crumb={urllib.parse.quote(YF_SESSION.crumb)}" if YF_SESSION.crumb else ""
        url = f"https://query1.finance.yahoo.com/v7/finance/quote?symbols={symbols_str}{crumb_param}"
        
        headers = STANDARD_HEADERS.copy()
        if YF_SESSION.cookie:
            headers['Cookie'] = YF_SESSION.cookie

        try:
            req = urllib.request.Request(url, headers=headers)
            with urllib.request.urlopen(req, timeout=12) as response:
                data = json.loads(response.read().decode('utf-8'))
                quotes = data.get('quoteResponse', {}).get('result', [])
                for q in quotes:
                    sym = q.get('symbol')
                    price = (q.get('regularMarketPrice') or 
                             q.get('regularMarketPreviousClose') or 
                             q.get('chartPreviousClose') or 
                             q.get('previousClose') or 0.0)
                    currency = q.get('currency', 'USD')
                    name = q.get('longName') or q.get('shortName') or sym
                    exchange = q.get('exchange', 'US')
                    mcap = q.get('marketCap', 0)
                    
                    curr_sym = "₹" if currency == "INR" or ".NS" in sym else "$"
                    if mcap >= 1e12:
                        market_cap = f"{curr_sym}{mcap / 1e12:.2f}T"
                    elif mcap >= 1e9:
                        market_cap = f"{curr_sym}{mcap / 1e9:.2f}B"
                    elif mcap > 0:
                        market_cap = f"{curr_sym}{mcap / 1e6:.2f}M"
                    else:
                        market_cap = "-"
                        
                    results_map[sym] = {
                        'name': name,
                        'exchange': exchange,
                        'market_cap': market_cap,
                        'ltp': price,
                        'currency': currency
                    }
        except Exception:
            pass
            
    missing_tickers = [t for t in tickers if t not in results_map or results_map[t]['ltp'] == 0.0]
    if missing_tickers:
        for sym in missing_tickers:
            try:
                crumb_str = f"&crumb={urllib.parse.quote(YF_SESSION.crumb)}" if YF_SESSION.crumb else ""
                sum_url = f"https://query2.finance.yahoo.com/v10/finance/quoteSummary/{sym}?modules=price,summaryDetail{crumb_str}"
                req = urllib.request.Request(sum_url, headers=headers)
                with urllib.request.urlopen(req, timeout=8) as resp:
                    sum_data = json.loads(resp.read().decode('utf-8'))
                    res_list = sum_data.get('quoteSummary', {}).get('result', [])
                    if res_list:
                        p_mod = res_list[0].get('price', {})
                        d_mod = res_list[0].get('summaryDetail', {})
                        price = (p_mod.get('regularMarketPrice', {}).get('raw') or 
                                 p_mod.get('regularMarketPreviousClose', {}).get('raw') or 0.0)
                        currency = p_mod.get('currency', 'USD')
                        name = p_mod.get('longName') or p_mod.get('shortName') or sym
                        exchange = p_mod.get('exchangeName') or 'US'
                        mcap = d_mod.get('marketCap', {}).get('raw') or p_mod.get('marketCap', {}).get('raw') or 0
                        
                        curr_sym = "₹" if currency == "INR" or ".NS" in sym else "$"
                        if mcap >= 1e12:
                            market_cap = f"{curr_sym}{mcap / 1e12:.2f}T"
                        elif mcap >= 1e9:
                            market_cap = f"{curr_sym}{mcap / 1e9:.2f}B"
                        elif mcap > 0:
                            market_cap = f"{curr_sym}{mcap / 1e6:.2f}M"
                        else:
                            market_cap = "-"
                            
                        results_map[sym] = {
                            'name': name,
                            'exchange': exchange,
                            'market_cap': market_cap,
                            'ltp': price,
                            'currency': currency
                        }
            except Exception:
                pass

    for t in tickers:
        if t not in results_map:
            results_map[t] = {
                'name': t,
                'exchange': 'US',
                'market_cap': '-',
                'ltp': 0.0,
                'currency': 'USD'
            }
            
    return results_map

def clean_sheet_name(name):
    invalid_chars = ['\\', '/', '?', '*', ':', '[', ']']
    for char in invalid_chars:
        name = name.replace(char, '_')
    return name[:31]

def run_etf_holdings_csv_lookup(holdings_csv_path, stock_codes_str, output_xlsx_path):
    start_time = time.time()
    
    abs_csv_path = os.path.abspath(holdings_csv_path)
    if not os.path.exists(abs_csv_path):
        print(f"[-] Error: ETF Holdings CSV file not found at: {abs_csv_path}")
        return

    print(f"[+] Loading ETF Holdings CSV database from: {abs_csv_path}")
    try:
        df_holdings = pd.read_csv(abs_csv_path)
    except Exception as e:
        print(f"[-] Error reading ETF Holdings CSV: {e}")
        return

    raw_stocks = [s.strip().upper() for s in stock_codes_str.split(",") if s.strip()]
    if not raw_stocks:
        print("[-] Error: No valid stock codes provided.")
        return

    print(f"[+] Querying metadata for input stocks: {raw_stocks}")
    stock_metadata = fetch_quotes_batch(raw_stocks)
    for stock in raw_stocks:
        meta = stock_metadata.get(stock, {})
        print(f"    -> {stock}: {meta.get('name', stock)} | Price: {meta.get('ltp', 0.0)} | Market Cap: {meta.get('market_cap', '-')}")

    # Gather matching ETFs across all matched ETFs for batch pricing
    all_matched_etfs = set()
    stock_to_matches = {}

    for stock in raw_stocks:
        stock_matches = df_holdings[df_holdings['Holding Ticker'].astype(str).str.strip().str.upper() == stock].copy()
        stock_matches['Weight_Numeric'] = pd.to_numeric(stock_matches['Holding Weight %'], errors='coerce').fillna(0.0)
        stock_matches = stock_matches.sort_values(by='Weight_Numeric', ascending=False)
        stock_to_matches[stock] = stock_matches
        for _, row in stock_matches.iterrows():
            all_matched_etfs.add(str(row['ETF Ticker']).strip().upper())

    sorted_all_etfs = list(all_matched_etfs)
    print(f"[+] Fetching live market prices for {len(sorted_all_etfs)} matching ETFs across all stocks...")
    etf_quotes = fetch_quotes_batch(sorted_all_etfs)

    # Generate Professional Excel Workbook (Exact V2 Layout)
    abs_out_path = os.path.abspath(output_xlsx_path)
    wb = openpyxl.Workbook()
    default_sheet = wb.active
    default_sheet.title = "ETF Holdings Analysis"

    navy_header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    alt_zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    light_tint_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    font_white_bold = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    font_regular = Font(name="Segoe UI", size=11)
    font_bold = Font(name="Segoe UI", size=11, bold=True)
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'), right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'), bottom=Side(style='thin', color='E0E0E0')
    )

    # 1. Populate Summary Sheet ("ETF Holdings Analysis") first
    ws_summary = default_sheet
    ws_summary.views.sheetView[0].showGridLines = True

    # Header Row 1: Stock Summary Information
    ws_summary.append(["Stock Summary Information"])
    ws_summary.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)
    ws_summary.cell(row=1, column=1).fill = navy_header_fill
    ws_summary.cell(row=1, column=1).font = font_white_bold
    ws_summary.cell(row=1, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.row_dimensions[1].height = 24

    # Header Row 2: Summary Table Columns
    summary_headers = ["Stock Ticker", "Company Name", "Exchange", "Market Cap", "Price", "ETF Codes (Comma separated)"]
    ws_summary.append(summary_headers)
    for col_idx in range(1, 7):
        cell = ws_summary.cell(row=2, column=col_idx)
        cell.fill = navy_header_fill
        cell.font = font_white_bold
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws_summary.row_dimensions[2].height = 26

    # Insert Summary Rows
    summary_row_idx = 3
    for stock in raw_stocks:
        meta = stock_metadata.get(stock, {'name': stock, 'exchange': 'US', 'market_cap': '-', 'ltp': 0.0, 'currency': 'USD'})
        stock_matches = stock_to_matches.get(stock, pd.DataFrame())
        etf_list = [str(r['ETF Ticker']).strip().upper() for _, r in stock_matches.iterrows()]
        etf_codes_str = ",".join(etf_list) if etf_list else "-"

        ws_summary.append([stock, meta['name'], meta['exchange'], meta['market_cap'], meta['ltp'], etf_codes_str])
        is_even = (summary_row_idx % 2 == 0)

        for col_idx in range(1, 7):
            cell = ws_summary.cell(row=summary_row_idx, column=col_idx)
            cell.font = font_regular
            cell.border = thin_border
            if is_even:
                cell.fill = alt_zebra_fill

            if col_idx in [1, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if col_idx == 1:
                    cell.font = font_bold
            elif col_idx == 2:
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 5:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                curr_sym = "₹" if meta['currency'] == 'INR' else "$"
                cell.number_format = f'"{curr_sym}"#,##0.00'
            elif col_idx == 6:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws_summary.row_dimensions[summary_row_idx].height = 22
        summary_row_idx += 1

    # Auto-fit columns for Summary Sheet
    for col in ws_summary.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 18)

    # 2. Populate individual stock sheets matching V2 format exactly
    for stock in raw_stocks:
        sheet_name = clean_sheet_name(stock)
        ws = wb.create_sheet(title=sheet_name)
        ws.views.sheetView[0].showGridLines = True

        # Row 1: Ticker in Column B
        ws.cell(row=1, column=2, value=stock)
        ws.cell(row=1, column=2).font = font_bold
        ws.row_dimensions[1].height = 22

        # Row 2: ETF Table Headers in Columns B to E (Col A blank or offset)
        etf_table_headers = ["ETF Ticker", "ETF Name", "ETF Price", "Holding Weight %"]
        for idx, h_text in enumerate(etf_table_headers, start=2):
            cell = ws.cell(row=2, column=idx, value=h_text)
            cell.fill = navy_header_fill
            cell.font = font_white_bold
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        ws.row_dimensions[2].height = 26

        stock_matches = stock_to_matches.get(stock, pd.DataFrame())
        current_io_row = 3

        if stock_matches.empty:
            ws.cell(row=current_io_row, column=2, value="-")
            ws.cell(row=current_io_row, column=3, value="No matching ETFs found")
            ws.cell(row=current_io_row, column=4, value="-")
            ws.cell(row=current_io_row, column=5, value=0.0)
            for col_idx in range(2, 6):
                cell = ws.cell(row=current_io_row, column=col_idx)
                cell.font = font_regular
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
            current_io_row += 1
        else:
            for _, row in stock_matches.iterrows():
                etf_ticker = str(row['ETF Ticker']).strip().upper()
                etf_name = str(row['ETF Name']).strip()
                
                try:
                    weight_val = float(row['Holding Weight %']) / 100.0
                except (ValueError, TypeError):
                    weight_val = 0.0
                    
                etf_meta = etf_quotes.get(etf_ticker, {'ltp': 0.0, 'currency': 'USD'})
                etf_price = etf_meta.get('ltp', 0.0)
                
                ws.cell(row=current_io_row, column=2, value=etf_ticker)
                ws.cell(row=current_io_row, column=3, value=etf_name)
                ws.cell(row=current_io_row, column=4, value=etf_price if etf_price > 0 else "-")
                ws.cell(row=current_io_row, column=5, value=weight_val)

                is_even = (current_io_row % 2 == 0)

                for col_idx in range(2, 6):
                    cell = ws.cell(row=current_io_row, column=col_idx)
                    cell.font = font_regular
                    cell.border = thin_border
                    if is_even:
                        cell.fill = alt_zebra_fill

                    if col_idx == 2:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        cell.font = font_bold
                    elif col_idx == 3:
                        cell.alignment = Alignment(horizontal="left", vertical="center")
                    elif col_idx == 4:
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                        if isinstance(cell.value, (int, float)):
                            curr_sym = "₹" if etf_meta.get('currency') == 'INR' else "$"
                            cell.number_format = f'"{curr_sym}"#,##0.00'
                    elif col_idx == 5:
                        cell.alignment = Alignment(horizontal="right", vertical="center")
                        cell.number_format = '0.00%'
                        cell.fill = light_tint_fill
                        cell.font = font_bold

                ws.row_dimensions[current_io_row].height = 22
                current_io_row += 1

        # Auto-fit column widths for individual sheet
        ws.column_dimensions['A'].width = 4
        for col in ws.iter_cols(min_col=2, max_col=5):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 18)

    wb.save(abs_out_path)
    
    elapsed = time.time() - start_time
    time_str = str(datetime.timedelta(seconds=int(elapsed)))

    print("\n" + "="*50)
    print("        ETF HOLDINGS LOOKUP REPORT EXECUTION")
    print("="*50)
    print(f"[>] Target Stock(s) Processed             : {raw_stocks}")
    print(f"[>] Total Execution Elapsed Time          : {time_str}")
    print(f"[>] Output Saved Location                 : {abs_out_path}")
    print("="*50 + "\n")

if __name__ == '__main__':
    parser = argparse.ArgumentParser(description="Reverse Stock-to-ETF Lookup Tool using ETF Holdings Detail CSV")
    parser.add_argument('--csv', type=str, default=DEFAULT_ETF_HOLDINGS_CSV, help="Path to ETF Holdings Detail CSV")
    parser.add_argument('--stocks', type=str, default=None, help="Target stock ticker(s) separated by comma")
    parser.add_argument('--output', type=str, default=DEFAULT_OUTPUT_XLSX, help="Output Excel filename")
    
    args = parser.parse_args()
    resolved_stocks = args.stocks if args.stocks is not None else FAIL_SAFE_STOCKCODE
    run_etf_holdings_csv_lookup(args.csv, resolved_stocks, args.output)