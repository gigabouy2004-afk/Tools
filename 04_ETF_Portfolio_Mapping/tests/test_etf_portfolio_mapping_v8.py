import importlib.util
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd


MODULE_PATH = Path(__file__).resolve().parents[1] / "ETF_Portfolio_Mapping_V8.py"
SPEC = importlib.util.spec_from_file_location("etf_portfolio_mapping_v8", MODULE_PATH)
B8 = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(B8)


class IssuerHoldingsTests(unittest.TestCase):
    def test_foto_combines_direct_and_receive_exposure(self):
        payload = {
            "holdings": [
                {"security_name": "RECV FOTO TRS LITE EQ", "security_ticker": None, "weight": 0.70},
                {"security_name": "Lumentum Holdings Inc", "security_ticker": "LITE", "weight": 0.80},
                {"security_name": "PAYB FOTO TRS LITE EQ", "security_ticker": None, "weight": -0.10},
                {"security_name": "CASH AND CASH EQUIVALENTS", "security_ticker": None, "weight": 98.60},
            ]
        }
        with patch.object(B8, "_read_url_text", return_value=json.dumps(payload)):
            rows = B8.fetch_holdings_from_foto_issuer("FOTO")

        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["Company Ticker"], "LITE")
        self.assertEqual(rows[0]["Company Name"], "Lumentum Holdings Inc")
        self.assertAlmostEqual(rows[0]["Weight"], 0.015)

    def test_corgi_percentage_points_below_one_are_not_magnified(self):
        holdings = [
            {
                "position_date": "2026-08-05",
                "security_name": "Taiwan Semiconductor",
                "security_cusip": "874039100",
                "security_ticker": "TSM",
                "weight_pct": "10.3700",
            },
            {
                "position_date": "2026-08-05",
                "security_name": "Tower Semiconductor",
                "security_cusip": "M87915274",
                "security_ticker": "TSEM",
                "weight_pct": "0.7600",
            },
        ]
        html = "<script>" + ",".join(json.dumps(item) for item in holdings) + "</script>"
        with patch.object(B8, "_read_url_text", return_value=html):
            rows = B8.fetch_holdings_from_corgi_issuer("EUV")

        weights = {row["Company Ticker"]: row["Weight"] for row in rows}
        self.assertAlmostEqual(weights["TSM"], 0.1037)
        self.assertAlmostEqual(weights["TSEM"], 0.0076)

    def test_tema_lazr_uses_full_issuer_rows_and_normalizes_exchange_codes(self):
        csv_text = """holdings_date,ticker,cusip,proper_name,shares,market_value,percent_of_nav,is_cash,country,sector
2026-08-04,LITE,55024U109,LUMENTUM HOLDINGS INC,3480,2956155.6,0.1517,0,United States,Information Technology
2026-08-04,AIXA GR,,AIXTRON SE,34512,1529745.6,0.0785,0,Germany,Information Technology
2026-08-04,3081 TT,,LANDMARK OPTOELECTRONICS CORP,12813.2,775003.86,0.0398,0,Taiwan,Information Technology
2026-08-04,ANTHROPIC SPV,ANTHROPIC SPV,ANTHROPIC SPV EXPOSURE,1820,2003820,0.1028,0,,
2026-08-04,,,CASH & CASH EQUIVALENTS,1,1,0.0112,1,,
"""
        with patch.object(B8, "_read_url_text", return_value=csv_text):
            rows = B8.fetch_holdings_from_tema_issuer("LAZR")

        by_ticker = {row["Company Ticker"]: row for row in rows}
        self.assertEqual(set(by_ticker), {"LITE", "AIXA.DE", "3081.TWO", "ANTHROPIC-SPV"})
        self.assertEqual(by_ticker["AIXA.DE"]["Source Ticker"], "AIXA GR")
        self.assertEqual(by_ticker["ANTHROPIC-SPV"]["Instrument Type"], "Private Company")
        self.assertAlmostEqual(by_ticker["LITE"]["Weight"], 0.1517)

    def test_tema_lazr_metrics_override_stale_yahoo_aum(self):
        html = """
        <span>AUM</span><div class="col-details">$19,484,972</div>
        <h2>LAZR NAV / Market Price</h2>
        <span>NAV</span><span>$42.36</span>
        <span>Market Price</span><span>$42.92</span>
        <div>As of August 04, 2026</div>
        """
        with patch.object(B8, "_read_url_text", return_value=html):
            metrics = B8.fetch_tema_fund_metrics("LAZR")

        self.assertEqual(metrics["Name"], "Tema Photonics & Optical ETF")
        self.assertEqual(metrics["AUM (USD M)"], 19.48)
        self.assertEqual(metrics["NAV"], 42.36)
        self.assertEqual(metrics["Market Price"], 42.92)

    def test_returns_do_not_relabel_since_inception_as_one_year(self):
        dates = pd.to_datetime(["2026-06-30", "2026-07-31", "2026-08-05"])
        prices = pd.Series([50.0, 40.0, 43.0], index=dates)
        performance = B8.calculate_price_performance(prices)

        self.assertIsNone(performance["3 Month (%)"])
        self.assertIsNone(performance["YTD (%)"])
        self.assertIsNone(performance["1 Year (%)"])


class InstrumentClassificationTests(unittest.TestCase):
    def test_lazr_issuer_identity_overrides_stale_yahoo_equity_type(self):
        instrument_type, basis = B8.classify_instrument(
            "LAZR",
            {
                "quoteType": "EQUITY",
                "longName": "Tema Photonics & Optical ETF",
            },
        )
        self.assertEqual(instrument_type, "ETF")
        self.assertEqual(basis, "Issuer-confirmed ETF ticker")

    def test_fund_name_overrides_equity_quote_type_for_reused_tickers(self):
        instrument_type, basis = B8.classify_instrument(
            "REUSED",
            {"quoteType": "EQUITY", "shortName": "Example Optical ETF"},
        )
        self.assertEqual(instrument_type, "ETF")
        self.assertEqual(basis, "Fund name")

    def test_nested_funds_are_not_reported_as_companies_or_stocks(self):
        holdings = pd.DataFrame(
            {
                "Company Ticker": ["LITE", "LAZR", "ANTHROPIC-SPV"],
                "Instrument Type": ["Stock", "ETF", "Private Company"],
            }
        )
        company_rows, stock_rows, nested_rows = B8.partition_holding_views(holdings)

        self.assertEqual(company_rows["Company Ticker"].tolist(), ["LITE", "ANTHROPIC-SPV"])
        self.assertEqual(stock_rows["Company Ticker"].tolist(), ["LITE"])
        self.assertEqual(nested_rows["Company Ticker"].tolist(), ["LAZR"])


class OutputContractTests(unittest.TestCase):
    def test_canonical_company_name_prevents_case_only_group_splits(self):
        selected = B8._select_canonical_company_name(
            ["LUMENTUM HOLDINGS INC", "Lumentum Holdings Inc"],
            "LITE",
        )
        self.assertEqual(selected, "Lumentum Holdings Inc")

    def test_etf_summary_sorts_by_mtd_descending(self):
        frame = pd.DataFrame(
            {
                "ETF_Code": ["EUV", "FOTO", "NONE"],
                "MTD (%)": [9.87, 15.78, None],
            }
        )
        result = B8.sort_etf_summary_by_mtd(frame)
        self.assertEqual(result["ETF_Code"].tolist(), ["FOTO", "EUV", "NONE"])

    def test_requested_weight_number_formats(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "formatting.xlsx"
            workbook = openpyxl.Workbook()
            etf = workbook.active
            etf.title = "ETF_Summary"
            etf.append([
                "ETF_Code", "Holding_Count", "Total_Holding_Weight", "Avg_Holding_Weight",
                "Max_Holding_Weight", "AUM (USD M)", "Last Trading Volume", "MTD (%)",
            ])
            etf.append(["FOTO", 15, 99.516, 6.6344, 15.515, 167.25, 278687, 15.78234])
            stock = workbook.create_sheet("Stock_Summary")
            stock.append([
                "Stock Ticker", "ETF_Count", "Total_Weight_Across_ETFs",
                "Avg_Weight_When_Present", "Max_Weight_In_One_ETF",
            ])
            stock.append(["LITE", 2, 19.585, 9.7925, 15.515])
            raw = workbook.create_sheet("RAW_Holdings")
            raw.append(["ETF_Code", "Stock Ticker", "Weight"])
            raw.append(["FOTO", "LITE", 0.15515])
            workbook.save(path)

            B8.apply_etf_summary_formatting(path)
            B8.apply_stock_summary_formatting(path)
            B8.apply_raw_holdings_formatting(path)

            checked = openpyxl.load_workbook(path)
            self.assertEqual(checked["ETF_Summary"]["C2"].number_format, "0.00")
            self.assertEqual(checked["ETF_Summary"]["H2"].number_format, "0.00")
            self.assertEqual(checked["Stock_Summary"]["C2"].number_format, "0.00")
            self.assertEqual(checked["RAW_Holdings"]["C2"].number_format, "0.00%")


if __name__ == "__main__":
    unittest.main()
