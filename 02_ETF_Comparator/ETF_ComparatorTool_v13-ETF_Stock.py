import argparse
import datetime as dt
from pathlib import Path

import openpyxl
import pandas as pd
import yfinance as yf
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill

# -----------------------------------------------------------------------------
# CONFIGURATION
# -----------------------------------------------------------------------------
BASE_FOLDER = Path("D:/Tools/ETF_Comparator")
INPUT_FILENAME = Path("D:/Tools/00_StockCodeMaster/02_Stock/0004-US-All_Energy_Codes.csv")
#INPUT_FILENAME = Path("D:/Tools/00_StockCodeMaster/03_ETF/17-07-US_ETF_Master_Library.csv")
OUTPUT_FILENAME = Path("D:/TMP/ETF_Comparator/20-07-2026-All_EnergyStocks_Investing-com_Comparator_Report.xlsx")

TICKER_COLUMN_CANDIDATES = (
    "Symbol", "Ticker", "Code", "Stock Code", "StockCode", "ACT Symbol",
    "NASDAQ Symbol", "CQS Symbol"
)
ETF_COLUMN_CANDIDATES = ("ETF", "ETF Flag", "Is ETF", "IsETF")
NAME_COLUMN_CANDIDATES = (
    "Security Name", "Name", "Company Name", "SecurityName"
)
EXCHANGE_COLUMN_CANDIDATES = (
    "Exchange", "Listing Exchange", "Market", "Market Category"
)


def clean_text(value):
    if value is None or pd.isna(value):
        return ""
    return str(value).strip()


def normalize_yes_no(value):
    text = clean_text(value).upper()
    if text in {"Y", "YES", "TRUE", "1", "ETF", "E"}:
        return "Y"
    if text in {"N", "NO", "FALSE", "0", "STOCK", "EQUITY", "S"}:
        return "N"
    return ""


def find_column(columns, candidates):
    lookup = {str(col).strip().lower(): col for col in columns}
    for candidate in candidates:
        found = lookup.get(candidate.lower())
        if found is not None:
            return found
    return None


def read_source_file(input_path):
    """Read CSV, TXT, pipe-delimited NASDAQ files, or Excel files."""
    suffix = input_path.suffix.lower()

    if suffix in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(input_path, dtype=str)

    # sep=None handles comma, tab and pipe-delimited NASDAQ Trader files.
    return pd.read_csv(
        input_path,
        dtype=str,
        encoding="utf-8-sig",
        sep=None,
        engine="python",
    )


def load_instruments(input_path):
    if not input_path.exists():
        raise FileNotFoundError(f"Input file not found: {input_path}")

    df = read_source_file(input_path)
    if df.empty:
        raise ValueError(f"Input file contains no rows: {input_path}")

    ticker_col = find_column(df.columns, TICKER_COLUMN_CANDIDATES)
    etf_col = find_column(df.columns, ETF_COLUMN_CANDIDATES)
    name_col = find_column(df.columns, NAME_COLUMN_CANDIDATES)
    exchange_col = find_column(df.columns, EXCHANGE_COLUMN_CANDIDATES)

    # Backward compatibility for a one-column ticker file.
    if ticker_col is None and len(df.columns) == 1:
        ticker_col = df.columns[0]

    if ticker_col is None:
        raise ValueError(
            "Ticker column not found. Expected one of: "
            + ", ".join(TICKER_COLUMN_CANDIDATES)
        )

    rows = []
    for _, source_row in df.iterrows():
        ticker = clean_text(source_row.get(ticker_col)).upper()
        if not ticker or ticker in {"SYMBOL", "TICKER", "CODE", "FILE CREATION TIME"}:
            continue

        etf_flag = normalize_yes_no(source_row.get(etf_col)) if etf_col else ""
        rows.append({
            "Ticker": ticker,
            "ETF Flag": etf_flag,
            "Source Name": clean_text(source_row.get(name_col)) if name_col else "",
            "Exchange": clean_text(source_row.get(exchange_col)) if exchange_col else "",
        })

    instruments = pd.DataFrame(rows)
    if instruments.empty:
        raise ValueError("No valid ticker rows were found in the input file.")

    instruments = instruments.drop_duplicates(subset=["Ticker"], keep="first")
    instruments = instruments.sort_values("Ticker").reset_index(drop=True)

    if etf_col is None:
        print("[WARN] No ETF column was found. Instrument type will be identified from Yahoo Finance.")
    else:
        known = instruments["ETF Flag"].isin(["Y", "N"]).sum()
        print(f"[INFO] ETF flag loaded for {known} of {len(instruments)} instruments.")

    return instruments


def parse_cli_tickers(ticker_text):
    if not ticker_text:
        return pd.DataFrame(columns=["Ticker", "ETF Flag", "Source Name", "Exchange"])

    tickers = sorted({x.strip().upper() for x in ticker_text.split(",") if x.strip()})
    return pd.DataFrame({
        "Ticker": tickers,
        "ETF Flag": "",
        "Source Name": "",
        "Exchange": "",
    })


def identify_instrument(ticker_obj, source_etf_flag):
    """Source ETF flag is authoritative; Yahoo is only a fallback."""
    if source_etf_flag == "Y":
        return "ETF", "Source ETF flag"
    if source_etf_flag == "N":
        return "Stock", "Source ETF flag"

    quote_type = ""
    try:
        quote_type = clean_text(ticker_obj.fast_info.get("quoteType")).upper()
    except Exception:
        pass

    if not quote_type:
        try:
            quote_type = clean_text(ticker_obj.info.get("quoteType")).upper()
        except Exception:
            pass

    if quote_type in {"ETF", "MUTUALFUND"}:
        return "ETF", "Yahoo quoteType fallback"
    if quote_type in {"EQUITY", "STOCK"}:
        return "Stock", "Yahoo quoteType fallback"

    # User's source filtering guarantees that non-ETF codes are valid stocks.
    return "Stock", "Default non-ETF"


def yahoo_symbol(ticker):
    value = ticker.strip().upper()
    if value.startswith("XNSE:"):
        return value.split(":", 1)[1] + ".NS"
    if value.startswith("XNSE") and ":" not in value:
        return value.replace("XNSE", "", 1) + ".NS"
    return value


def fetch_price_history(ticker):
    """
    Fetch split/dividend-adjusted prices explicitly.

    auto_adjust=True is deliberate for both ETFs and stocks so a normal stock
    split does not appear as a price gain or loss. The same adjusted series is
    used for every return period.
    """
    yf_symbol = yahoo_symbol(ticker)
    ticker_obj = yf.Ticker(yf_symbol)
    history = ticker_obj.history(
        period="max",
        auto_adjust=True,
        actions=True,
        repair=True,
    )

    if history is None or history.empty or "Close" not in history.columns:
        return ticker_obj, None, None

    history.index = pd.to_datetime(history.index).tz_localize(None)
    prices = pd.to_numeric(history["Close"], errors="coerce").dropna().sort_index()
    if prices.empty:
        return ticker_obj, None, history

    return ticker_obj, prices, history



def get_info(ticker_obj):
    try:
        return ticker_obj.info or {}
    except Exception:
        return {}


def get_size_value_usd_m(ticker_obj, info, instrument_type):
    """Return ETF AUM or stock market capitalization in USD millions."""
    value = None

    if instrument_type == "ETF":
        value = info.get("totalAssets")
    else:
        value = info.get("marketCap")
        if value is None:
            try:
                value = ticker_obj.fast_info.get("marketCap")
            except Exception:
                value = None

    try:
        return round(float(value) / 1_000_000.0, 2) if value is not None else None
    except (TypeError, ValueError):
        return None


def latest_trading_volume(history):
    """Return actual reported trading volume for the latest available session."""
    if history is None or history.empty or "Volume" not in history.columns:
        return None

    volume = pd.to_numeric(history["Volume"], errors="coerce").dropna()
    if volume.empty:
        return None

    return int(volume.iloc[-1])


def compute_risk_metrics(prices, benchmark_prices, instrument_type, risk_free_rate=0.04):
    """Calculate beta for both types and alpha only for ETFs."""
    result = {"Beta": None, "Alpha (Ann. %)": None}
    if prices is None or benchmark_prices is None:
        return result

    aligned = pd.concat(
        [prices.rename("Instrument"), benchmark_prices.rename("Benchmark")],
        axis=1,
        join="inner",
    ).dropna()

    if len(aligned) < 30:
        return result

    aligned = aligned.iloc[-756:]
    returns = aligned.pct_change().dropna()
    if len(returns) < 20:
        return result

    instrument_returns = returns["Instrument"]
    benchmark_returns = returns["Benchmark"]
    benchmark_variance = benchmark_returns.var(ddof=1)
    if pd.isna(benchmark_variance) or benchmark_variance == 0:
        return result

    beta = instrument_returns.cov(benchmark_returns) / benchmark_variance
    result["Beta"] = round(float(beta), 2)

    if instrument_type == "ETF":
        daily_rf = risk_free_rate / 252.0
        alpha_daily = (instrument_returns - daily_rf).mean() - beta * (benchmark_returns - daily_rf).mean()
        result["Alpha (Ann. %)"] = round(float(alpha_daily * 252.0 * 100.0), 2)

    return result

def calculate_performance(prices):
    if prices is None or len(prices) < 2:
        return {}

    prices = prices.dropna().sort_index()
    last_date = pd.Timestamp(prices.index[-1])

    def index_on_or_before(date_value):
        position = prices.index.searchsorted(pd.Timestamp(date_value), side="right") - 1
        return position if position >= 0 else None

    def index_before(date_value):
        position = prices.index.searchsorted(pd.Timestamp(date_value), side="left") - 1
        return position if position >= 0 else None

    def index_on_or_after(date_value):
        position = prices.index.searchsorted(pd.Timestamp(date_value), side="left")
        return position if position < len(prices) else None

    def return_between(start_date, end_date=None, use_previous_close=False):
        end_date = last_date if end_date is None else pd.Timestamp(end_date)

        if use_previous_close:
            start_index = index_before(start_date)
        else:
            start_index = index_on_or_before(start_date)

        if start_index is None:
            start_index = index_on_or_after(start_date)

        end_index = index_on_or_before(end_date)
        if start_index is None or end_index is None or end_index <= start_index:
            return None

        start_price = float(prices.iloc[start_index])
        end_price = float(prices.iloc[end_index])
        if start_price <= 0:
            return None

        # This formula preserves the sign: lower end price = negative return.
        return round(((end_price / start_price) - 1.0) * 100.0, 2)

    week_start = last_date.normalize() - pd.Timedelta(days=last_date.weekday())
    month_start = pd.Timestamp(last_date.year, last_date.month, 1)
    year_start = pd.Timestamp(last_date.year, 1, 1)

    previous_month_end = month_start - pd.Timedelta(days=1)
    previous_month_start = pd.Timestamp(previous_month_end.year, previous_month_end.month, 1)

    two_months_ago_end = previous_month_start - pd.Timedelta(days=1)
    two_months_ago_start = pd.Timestamp(two_months_ago_end.year, two_months_ago_end.month, 1)

    result = {
        "Price": round(float(prices.iloc[-1]), 2),
        "Since Yesterday (%)": round(
            ((float(prices.iloc[-1]) / float(prices.iloc[-2])) - 1.0) * 100.0, 2
        ),
        "This Week (%)": return_between(week_start, use_previous_close=True),
        "MTD (%)": return_between(month_start, use_previous_close=True),
        previous_month_start.strftime("%b-%y (%%)"): return_between(
            previous_month_start, previous_month_end, use_previous_close=True
        ),
        two_months_ago_start.strftime("%b-%y (%%)"): return_between(
            two_months_ago_start, two_months_ago_end, use_previous_close=True
        ),
        "3 Month (%)": return_between(last_date - pd.DateOffset(months=3)),
        "YTD (%)": return_between(year_start, use_previous_close=True),
        "6 Month (%)": return_between(last_date - pd.DateOffset(months=6)),
        "9 Month (%)": return_between(last_date - pd.DateOffset(months=9)),
        "1 Year (%)": return_between(last_date - pd.DateOffset(years=1)),
    }
    return result


def recent_split_text(history, last_date):
    if history is None or history.empty or "Stock Splits" not in history.columns:
        return ""

    split_series = pd.to_numeric(history["Stock Splits"], errors="coerce").fillna(0)
    split_series = split_series[split_series != 0]
    if split_series.empty:
        return ""

    cutoff = pd.Timestamp(last_date) - pd.DateOffset(years=1)
    recent = split_series[split_series.index >= cutoff]
    if recent.empty:
        return ""

    return "; ".join(
        f"{pd.Timestamp(index).date()} x{value:g}" for index, value in recent.items()
    )


def best_name(info, source_name):
    if source_name:
        return source_name
    return clean_text(info.get("longName") or info.get("shortName"))


def apply_excel_formatting(filepath, dataframe):
    """Apply the original V12 three-rank layout and match-based cell colours."""
    workbook = openpyxl.load_workbook(filepath)
    worksheet = workbook.active
    worksheet.title = "Price Performance"

    worksheet.insert_rows(1, 3)
    column_numbers = {name: i + 1 for i, name in enumerate(dataframe.columns)}
    percentage_columns = [name for name in dataframe.columns if name.endswith("(%)")]

    # Original visual ranking design for every performance-period column:
    # row 1 = maximum, row 2 = second-highest DISTINCT value, row 3 = minimum.
    # The reference cells themselves carry fixed red/yellow/green fills.
    red_fill = PatternFill(fill_type="solid", fgColor="FFFF0000")
    yellow_fill = PatternFill(fill_type="solid", fgColor="FFFFFF00")
    green_fill = PatternFill(fill_type="solid", fgColor="FF92D050")

    first_data_row = 5
    last_data_row = len(dataframe) + 4

    for column_name in percentage_columns:
        values = pd.to_numeric(dataframe[column_name], errors="coerce").dropna()
        unique_desc = sorted({round(float(v), 2) for v in values}, reverse=True)

        maximum = unique_desc[0] if unique_desc else None
        second_maximum = unique_desc[1] if len(unique_desc) >= 2 else None
        minimum = unique_desc[-1] if unique_desc else None
        reference_values = (maximum, second_maximum, minimum)
        reference_fills = (red_fill, yellow_fill, green_fill)

        column_number = column_numbers[column_name]
        letter = openpyxl.utils.get_column_letter(column_number)

        # Populate and permanently colour the three reference cells.
        for row_number, (value, fill) in enumerate(
            zip(reference_values, reference_fills), start=1
        ):
            cell = worksheet.cell(row=row_number, column=column_number, value=value)
            cell.font = Font(bold=True)
            cell.fill = fill
            cell.number_format = "0.00"

        data_range = f"{letter}{first_data_row}:{letter}{last_data_row}"
        # Keep the original conditional-formatting rules so that Excel can
        # re-evaluate the colours if values are edited after generation.
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="equal", formula=[f"{letter}$1"], fill=red_fill),
        )
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="equal", formula=[f"{letter}$2"], fill=yellow_fill),
        )
        worksheet.conditional_formatting.add(
            data_range,
            CellIsRule(operator="equal", formula=[f"{letter}$3"], fill=green_fill),
        )

        # Apply the same fills directly as well. Some spreadsheet previewers
        # do not evaluate conditional formatting until the file is opened and
        # recalculated in desktop Excel. Direct fills make the V12 ranking
        # colours visible immediately, while the conditional rules remain.
        rank_fills = {}
        if maximum is not None:
            rank_fills[maximum] = red_fill
        if second_maximum is not None:
            rank_fills[second_maximum] = yellow_fill
        if minimum is not None:
            rank_fills[minimum] = green_fill

        for row_number in range(first_data_row, last_data_row + 1):
            data_cell = worksheet.cell(row=row_number, column=column_number)
            if data_cell.value is None:
                continue
            try:
                rounded_value = round(float(data_cell.value), 2)
            except (TypeError, ValueError):
                continue
            matched_fill = rank_fills.get(rounded_value)
            if matched_fill is not None:
                data_cell.fill = matched_fill

    # V12-style number presentation.
    for row in range(5, worksheet.max_row + 1):
        for name in percentage_columns:
            worksheet.cell(row=row, column=column_numbers[name]).number_format = "0.00"

        for name in ("Price", "Beta", "Alpha (Ann. %)"):
            if name in column_numbers:
                worksheet.cell(row=row, column=column_numbers[name]).number_format = "0.00"

        for name in ("AUM / Market Cap (USD M)", "Last Trading Volume"):
            if name in column_numbers:
                worksheet.cell(row=row, column=column_numbers[name]).number_format = "#,##0"

    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = (
        f"A4:{openpyxl.utils.get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    )

    for column_cells in worksheet.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        worksheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(max_length + 2, 10), 45
        )

    workbook.save(filepath)


def final_output_path(output_path, cli_used):
    if not cli_used:
        return output_path
    timestamp = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    return output_path.with_name(f"{output_path.stem}-{timestamp}-CLI{output_path.suffix}")


def run(input_path, output_path, ticker_text=None):
    instruments = parse_cli_tickers(ticker_text) if ticker_text else load_instruments(input_path)
    output_path = final_output_path(output_path, bool(ticker_text))
    output_path.parent.mkdir(parents=True, exist_ok=True)

    results = []
    skipped = []

    print(f"[INFO] Instruments selected: {len(instruments)}")
    print("[INFO] Fetching SPY benchmark for Beta / Alpha calculation...")
    _, benchmark_prices, _ = fetch_price_history("SPY")

    for count, source in enumerate(instruments.to_dict("records"), start=1):
        ticker = source["Ticker"]
        try:
            ticker_obj, prices, history = fetch_price_history(ticker)
            if prices is None or len(prices) < 2:
                skipped.append((ticker, "No usable price history"))
                print(f"[SKIP] {ticker}: no usable price history")
                continue

            instrument_type, _ = identify_instrument(ticker_obj, source["ETF Flag"])
            info = get_info(ticker_obj)
            performance = calculate_performance(prices)
            risk_metrics = compute_risk_metrics(prices, benchmark_prices, instrument_type)

            row = {
                "Ticker": ticker,
                "ETF Flag": "Y" if instrument_type == "ETF" else "N",
                "Name": best_name(info, source["Source Name"]),
                "AUM / Market Cap (USD M)": get_size_value_usd_m(
                    ticker_obj, info, instrument_type
                ),
                "Last Trading Volume": latest_trading_volume(history),
                "Price": performance.pop("Price", None),
                **risk_metrics,
                **performance,
            }
            results.append(row)
            print(f"[OK {count}/{len(instruments)}] {ticker} -> {instrument_type}")

        except Exception as exc:
            skipped.append((ticker, str(exc)))
            print(f"[ERROR] {ticker}: {exc}")

    if not results:
        raise RuntimeError("No instruments produced a valid price-performance result.")

    dataframe = pd.DataFrame(results)

    fixed_front = [
        "Ticker", "ETF Flag", "Name", "AUM / Market Cap (USD M)",
        "Last Trading Volume", "Price", "Beta", "Alpha (Ann. %)",
        "Since Yesterday (%)", "This Week (%)", "MTD (%)",
    ]
    fixed_back = [
        "3 Month (%)", "YTD (%)", "6 Month (%)", "9 Month (%)",
        "1 Year (%)",
    ]
    middle = [c for c in dataframe.columns if c not in fixed_front + fixed_back]
    dataframe = dataframe.reindex(columns=fixed_front + middle + fixed_back)

    dataframe["MTD (%)"] = pd.to_numeric(dataframe["MTD (%)"], errors="coerce")
    dataframe = dataframe.sort_values(
        by=["MTD (%)", "Ticker"],
        ascending=[False, True],
        na_position="last",
    ).reset_index(drop=True)

    dataframe.to_excel(output_path, index=False, engine="openpyxl")
    apply_excel_formatting(output_path, dataframe)

    print("\n" + "=" * 88)
    print("STOCK / ETF PRICE COMPARATOR - POST EXECUTION SUMMARY")
    print("=" * 88)
    print(f"Input file          : {input_path if not ticker_text else 'CLI tickers'}")
    print(f"Codes selected      : {len(instruments)}")
    print(f"Codes processed     : {len(dataframe)}")
    print(f"ETF                 : {(dataframe['ETF Flag'] == 'Y').sum()}")
    print(f"Stock               : {(dataframe['ETF Flag'] == 'N').sum()}")
    print(f"Error / skipped     : {len(skipped)}")
    print("Sort order          : MTD (%) descending")
    print(f"Output file         : {output_path}")
    print("=" * 88)

    return output_path


def parse_args():
    parser = argparse.ArgumentParser(
        description="Identify stock/ETF codes and compare adjusted price performance."
    )
    parser.add_argument(
        "--input", type=Path, default=INPUT_FILENAME,
        help=f"NASDAQ/NYSE master CSV/TXT/XLSX containing Symbol and ETF flag. Default: {INPUT_FILENAME}",
    )
    parser.add_argument(
        "--output", type=Path, default=OUTPUT_FILENAME,
        help=f"Output XLSX file. Default: {OUTPUT_FILENAME}",
    )
    parser.add_argument(
        "--tickers", default=None,
        help="Optional comma-separated codes. Yahoo quoteType is used because no source ETF flag is available.",
    )
    return parser.parse_args()


if __name__ == "__main__":
    arguments = parse_args()
    run(arguments.input, arguments.output, arguments.tickers)