import argparse
import datetime
import pandas as pd
import yfinance as yf
from pathlib import Path
import openpyxl
from openpyxl.formatting.rule import ColorScaleRule

# --- CONFIG ---
BASE_FOLDER = Path("D:/Tools/ETF_Comparator")

#INPUT_FILENAME = Path("")
INPUT_FILENAME = Path("D:\\Tools\\StockCodeMaster\\03_ETF\\01-07-US_ETF_Master_Library.csv")  # Default input file for tickers
#INPUT_FILENAME = Path("D:\\Tools\\StockCodeMaster\\03_ETF\\01-07-US_ETF_Master_Library-China.csv")  # Default input file for tickers

#INPUT_FILENAME = None   
OUTPUT_FILENAME = Path("D:\\TMP\\07-07-ALL_ETF-Comparison-Report.xlsx")
CLASSIFICATION_FILENAME = Path("D:/Tools/StockCodeMaster/03_ETF/01-07-US_ETF_Classification_Mapping.csv")
print("***************************************************************************")
print("ETF Input file name - ", INPUT_FILENAME,"\n")
print("ETF Comparison file name - ", OUTPUT_FILENAME,"\n")
print("ETF Classification file name - ", CLASSIFICATION_FILENAME,"\n")


def parse_ticker_codes(ticker_text):
    if not ticker_text:
        return []

    tickers = [
        ticker.strip().upper()
        for ticker in ticker_text.split(",")
        if ticker.strip()
    ]
    return sorted(set(tickers))


def load_tickers(input_path, required=False):
    if not input_path.exists():
        if required:
            raise FileNotFoundError(f"Missing input file: {input_path}")
        print(f"[WARN] Input file not found: {input_path}")
        return []

    df = pd.read_csv(input_path, header=None, usecols=[0], dtype=str, encoding='utf-8-sig')
    tickers = (
        df.iloc[:, 0]
        .dropna()
        .astype(str)
        .str.strip()
        .str.upper()
        .unique()
        .tolist()
    )
    header_values = {"TICKER", "TICKERS", "SYMBOL", "SYMBOLS", "STOCK", "STOCK CODE", "STOCKCODE", "CODE"}
    tickers = [ticker for ticker in tickers if ticker and ticker not in header_values]
    return sorted(tickers)


def collect_tickers(input_path, ticker_text=None):
    cli_tickers = parse_ticker_codes(ticker_text)
    
    if cli_tickers:
        tickers = sorted(set(cli_tickers))
        print(f"[INFO] Using CLI tickers: {len(tickers)} (File input completely bypassed)")
    else:
        file_tickers = load_tickers(input_path)
        tickers = sorted(set(file_tickers))
        print(f"[INFO] Using File tickers: {len(tickers)} (No CLI input provided, reading from file)")

    if not tickers:
        raise ValueError("No ETF tickers found. Provide an input file or --tickers SMH,SOXX.")

    return tickers


def load_classification_mapping(mapping_path):
    if mapping_path.exists():
        return pd.read_csv(mapping_path, dtype=str).fillna("")

    raise FileNotFoundError(
        f"Classification mapping not found: {mapping_path}. "
        "Build it first with ETF_ClassificationMapper.py."
    )


def filter_tickers_by_classification(tickers, mapping_df, classification):
    if not classification:
        return tickers

    terms = [term.strip().lower() for term in classification.split(",") if term.strip()]
    searchable_cols = ["Asset Class", "Strategy", "Theme", "Geography", "Category Flags", "Security Name"]

    mask = pd.Series(False, index=mapping_df.index)
    for term in terms:
        term_mask = pd.Series(False, index=mapping_df.index)
        for col in searchable_cols:
            term_mask = term_mask | mapping_df[col].str.lower().str.contains(term, regex=False, na=False)
        mask = mask | term_mask

    selected = sorted(set(mapping_df.loc[mask, "Ticker"].str.upper()) & set(tickers))
    print(f"[INFO] Classification '{classification}' selected {len(selected)} tickers from {len(tickers)} available tickers")
    if not selected:
        raise ValueError(f"No tickers matched classification: {classification}")
    return selected


def get_price_series(ticker):
    if ticker.upper().startswith("XNSE"):
        ticker = ticker.upper().replace("XNSE:", "").replace("XNSE", "") + ".NS"

    tk = yf.Ticker(ticker)
    hist = tk.history(period="max")

    if hist.empty:
        return None, None, None

    hist.index = hist.index.tz_localize(None)

    price_col = "Adj Close" if "Adj Close" in hist.columns else "Close"
    prices = hist[price_col].dropna()

    return tk, prices, price_col


def compute_returns(prices: pd.Series):

    if prices is None or prices.empty or len(prices) < 2:
        return {}

    prices = prices.dropna().sort_index()
    last_date = prices.index[-1]

    def idx_on_or_before(dt):
        pos = prices.index.searchsorted(pd.Timestamp(dt), side="right") - 1
        return pos if pos >= 0 else None

    def idx_before(dt):
        pos = prices.index.searchsorted(pd.Timestamp(dt), side="left") - 1
        return pos if pos >= 0 else None

    def idx_on_or_after(dt):
        pos = prices.index.searchsorted(pd.Timestamp(dt), side="left")
        return pos if pos < len(prices) else None

    def calc_from_indices(s_idx, e_idx):
        if s_idx is None or e_idx is None or e_idx <= s_idx:
            return None

        start_price = prices.iloc[s_idx]
        end_price = prices.iloc[e_idx]
        if pd.isna(start_price) or pd.isna(end_price) or start_price == 0:
            return None

        return (end_price / start_price) - 1

    def calc_return(start_dt, end_dt=None, start_from_previous_close=False):
        """Return between two dates, optionally using the close before start_dt."""
        if end_dt is None:
            end_dt = last_date

        if start_from_previous_close:
            s_idx = idx_before(start_dt)
            if s_idx is None:
                s_idx = idx_on_or_after(start_dt)
        else:
            s_idx = idx_on_or_before(start_dt)
            if s_idx is None:
                s_idx = idx_on_or_after(start_dt)

        e_idx = idx_on_or_before(end_dt)
        return calc_from_indices(s_idx, e_idx)

    def pct(val):
        return round(val * 100, 2) if val is not None else None

    results = {}

    # --- Since Yesterday ---
    results["Since Yesterday (%)"] = (
        round(((prices.iloc[-1] / prices.iloc[-2]) - 1) * 100, 2)
        if len(prices) >= 2 else None
    )

    # --- WTD ---
    week_start = last_date - pd.Timedelta(days=last_date.weekday())
    results["This Week (%)"] = pct(calc_return(week_start, last_date, start_from_previous_close=True))

    # --- MTD ---
    month_start = pd.Timestamp(last_date.year, last_date.month, 1)
    results["MTD (%)"] = pct(calc_return(month_start, last_date, start_from_previous_close=True))

    # --- Fixed months ---
    months = {
        "May-26 (%)": (pd.Timestamp(2026, 5, 1), pd.Timestamp(2026, 5, 31)),
        "June-26 (%)": (pd.Timestamp(2026, 6, 1), pd.Timestamp(2026, 6, 30)),
    }

    for k, (s, e) in months.items():
        results[k] = pct(calc_return(s, e, start_from_previous_close=True))

    # --- YTD ---
    year_start = pd.Timestamp(last_date.year, 1, 1)
    results["YTD (%)"] = pct(calc_return(year_start, start_from_previous_close=True))

    # --- Rolling / trailing returns ---
    rolling_periods = {
        "3 Month (%)": pd.Timedelta(days=90),
        "6 Month (%)": pd.DateOffset(months=6),
        "9 Month (%)": pd.DateOffset(months=9),
        "1 yr (%)": pd.DateOffset(years=1),
        "3 yr (%)": pd.DateOffset(years=3),
    }

    for label, lookback in rolling_periods.items():
        start_dt = last_date - lookback
        results[label] = pct(calc_return(start_dt))

    return results


def normalize_ratio(value):
    if value is None or pd.isna(value):
        return None

    if isinstance(value, str):
        cleaned = value.strip().replace("%", "")
        if not cleaned:
            return None
        try:
            value = float(cleaned)
        except ValueError:
            return None

    try:
        value = float(value)
    except (TypeError, ValueError):
        return None

    if value > 1:
        value = value / 100

    return round(value, 2)


def get_expense_ratio_from_funds_data(tk):
    try:
        funds_data = tk.get_funds_data()
        fund_operations = funds_data.fund_operations
    except Exception:
        return None

    if fund_operations is None or fund_operations.empty:
        return None

    possible_labels = [
        "Annual Report Expense Ratio",
        "Expense Ratio",
        "Net Expense Ratio",
        "Gross Expense Ratio",
    ]

    for label in possible_labels:
        if label not in fund_operations.index:
            continue

        row = fund_operations.loc[label]
        if isinstance(row, pd.Series):
            for value in row.dropna():
                ratio = normalize_ratio(value)
                if ratio is not None:
                    return ratio
        else:
            ratio = normalize_ratio(row)
            if ratio is not None:
                return ratio

    return None


def get_expense_ratio(info, tk):
    possible_keys = [
        "expenseRatio",
        "annualReportExpenseRatio",
        "netExpenseRatio",
        "grossExpenseRatio",
    ]

    for key in possible_keys:
        ratio = normalize_ratio(info.get(key))
        if ratio is not None:
            return ratio

    return get_expense_ratio_from_funds_data(tk)


def extract_fundamentals(tk):
    try:
        info = tk.info
    except Exception:
        info = {}

    expense_ratio = get_expense_ratio(info, tk)

    aum = info.get("totalAssets") or info.get("marketCap")
    aum_m = round((aum / 1_000_000), 2) if aum else None

    return {
        "Name": info.get("longName", None),
        "AUM (USD M)": aum_m,
        "Expense Ratio": expense_ratio,
        "Liquidity (Avg Vol)": info.get("averageVolume", None),
    }


def get_final_output_path(output_path, ticker_text, classification, mapping_provided_via_cli, output_provided_via_cli):
    stem = output_path.stem
    cli_parts = []
    
    if mapping_provided_via_cli:
        cli_parts.append("Mappingfile")
        
    if classification:
        cli_parts.append("Classification")
        
    if ticker_text:
        cli_parts.append("ETFCodes")

    if output_provided_via_cli:
        cli_parts.append("Output")

    if cli_parts:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        stem += f"-{timestamp}-CLI-" + "-".join(cli_parts)
        
    return output_path.with_name(f"{stem}{output_path.suffix}")


def apply_excel_formatting(filepath, df):
    wb = openpyxl.load_workbook(str(filepath))
    ws = wb.active

    ws.insert_rows(1, 3)

    pct_cols = [col for col in df.columns if "(%)" in col]
    col_indices = {col: df.columns.get_loc(col) + 1 for col in df.columns}

    for col in pct_cols:
        col_idx = col_indices[col]
        col_data = pd.to_numeric(df[col], errors='coerce').dropna()
        
        if not col_data.empty:
            ws.cell(row=1, column=col_idx, value=round(col_data.max(), 2))
            ws.cell(row=2, column=col_idx, value=round(col_data.mean(), 2))
            ws.cell(row=3, column=col_idx, value=round(col_data.min(), 2))
            
            for r in range(1, 4):
                ws.cell(row=r, column=col_idx).font = openpyxl.styles.Font(bold=True)

    ws.cell(row=1, column=1, value="MAX Spread").font = openpyxl.styles.Font(bold=True)
    ws.cell(row=2, column=1, value="AVG Spread").font = openpyxl.styles.Font(bold=True)
    ws.cell(row=3, column=1, value="MIN Spread").font = openpyxl.styles.Font(bold=True)

    ws.freeze_panes = "A5"

    max_row = len(df) + 4
    for col in pct_cols:
        col_letter = openpyxl.utils.get_column_letter(col_indices[col])
        cell_range = f"{col_letter}5:{col_letter}{max_row}"
        
        rule = ColorScaleRule(
            start_type="min", start_color="FF9999",
            mid_type="percentile", mid_value=50, mid_color="FFFF99",
            end_type="max", end_color="99FF99"
        )
        ws.conditional_formatting.add(cell_range, rule)

    wb.save(str(filepath))


def run(ticker_text=None, classification=None, mapping_path=None, output_path_arg=None):
    base = Path(BASE_FOLDER)

    if mapping_path:
        input_path = mapping_path if mapping_path.is_absolute() else base / mapping_path
    else:
        input_path = INPUT_FILENAME if INPUT_FILENAME.is_absolute() else base / INPUT_FILENAME

    if output_path_arg:
        output_path = output_path_arg if output_path_arg.is_absolute() else base / output_path_arg
    else:
        output_path = OUTPUT_FILENAME if OUTPUT_FILENAME.is_absolute() else base / OUTPUT_FILENAME
    
    mapping_provided_via_cli = mapping_path is not None
    output_provided_via_cli = output_path_arg is not None
    actual_mapping_path = mapping_path if mapping_path else CLASSIFICATION_FILENAME
    
    output_path = get_final_output_path(
        output_path,
        ticker_text,
        classification,
        mapping_provided_via_cli,
        output_provided_via_cli,
    )
    
    if not output_path.parent.exists():
        raise FileNotFoundError(f"Output directory does not exist: {output_path.parent}. Aborting operation.")

    tickers = collect_tickers(input_path, ticker_text)

    if classification:
        mapping_df = load_classification_mapping(actual_mapping_path)
        tickers = filter_tickers_by_classification(tickers, mapping_df, classification)

    results = []

    for t in tickers:
        try:
            tk, prices, price_col = get_price_series(t)

            if prices is None:
                print(f"[SKIP] No data: {t}")
                continue

            current_price = round(prices.iloc[-1],2)

            returns = compute_returns(prices)
            fundamentals = extract_fundamentals(tk)

            row = {
                "Ticker": t,
                "Price": current_price,
                **fundamentals,
                **returns,
            }

            results.append(row)
            print(f"[OK] {t}")

        except Exception as e:
            print(f"[ERR] {t}: {e}")

    df = pd.DataFrame(results)

    ordered_cols = [
        "Ticker", "Name", "AUM (USD M)", "Price", "Liquidity (Avg Vol)",
        "Since Yesterday (%)", "This Week (%)", "MTD (%)",
        "June-26 (%)", "May-26 (%)", 
        "3 Month (%)", "YTD (%)",  "6 Month (%)", "9 Month (%)", "1 yr (%)", "3 yr (%)"

    ]

    df = df.reindex(columns=ordered_cols)

    if "MTD (%)" in df.columns:
        df["MTD (%)"] = pd.to_numeric(df["MTD (%)"], errors="coerce")
        df = df.sort_values(by=["MTD (%)", "June-26 (%)"], ascending=[False, False])

    try:
        df.to_excel(output_path, index=False, engine="openpyxl")
        apply_excel_formatting(output_path, df)
        print(f"\nSaved -> {output_path}")
    except PermissionError:
        tmp_output_path = output_path.with_name(f"{output_path.stem}-TMP{output_path.suffix}")
        df.to_excel(tmp_output_path, index=False, engine="openpyxl")
        apply_excel_formatting(tmp_output_path, df)
        print(f"\n[WARN] Write access denied for original file. Saved to -> {tmp_output_path}")


def parse_args():
    parser = argparse.ArgumentParser(description="Compare ETF returns and fundamentals.")
    parser.add_argument(
        "--tickers",
        help="Optional comma-separated ETF tickers to override INPUT_FILENAME completely, e.g. SMH,SOXX,XLK.",
    )
    parser.add_argument(
        "--classification",
        help="Optional comma-separated classification filter, e.g. Semiconductor, Bitcoin, Municipal Bond, India.",
    )
    parser.add_argument(
        "--classification-mapping-file",
        type=Path,
        default=None,
        help=f"Classification mapping CSV path. Default: {CLASSIFICATION_FILENAME}",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help=f"Output XLSX path. Default: {OUTPUT_FILENAME}",
    )
    return parser.parse_args()


if __name__ == "__main__":
    args = parse_args()
    run(args.tickers, args.classification, args.classification_mapping_file, args.output)
